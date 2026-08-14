# detection_entry_api.py
import requests
import json
import time
import base64
import mimetypes
from datetime import datetime, timedelta
import urllib3
from urllib3.exceptions import InsecureRequestWarning
import pandas as pd
import os
import re
import unicodedata
import paths
from openpyxl import load_workbook

# 禁用SSL警告
urllib3.disable_warnings(InsecureRequestWarning)


def _detection_no_of(code):
    """从样品号提取报验编号(字母前缀+8位流水)，作为服务端 keyword 查询值。
    输入可能带 3 位小号(001)、平行字母(A/B)或非标准后缀(如 _报告)，
    统一截取开头的 字母+8位；不符该结构(如已是无小号的报验编号)原样返回。"""
    m = re.match(r'[A-Za-z]+\d{8}', (code or '').strip())
    return m.group(0) if m else (code or '')

class DetectionAPI:
    def __init__(self, login_system):
        self.login_system = login_system
        # 缓存配置数据
        self.cached_units = None
        self.cached_round_methods = None
        self.cached_calc_methods = None
        self.current_sample_id = None
        self.sub_method_map = {}
        self.method_id_to_standard_no = {}
        self._std_no_name_to_id = self._load_std_no_name_cache()  # standardNoName -> methodId，持久化跨进程复用
        self.cached_solution_types = None

        # 实验编号缓存
        self.experiment_code_cache = {}  # 缓存已生成的实验编号
        self.user_prefix_cache = {}  # 缓存用户前缀

        self.load_sub_method_config()

    def generate_experiment_code(self, method_name="", log_func=None, force_new=False):
        """生成实验编号 - 统一实现，避免重复。
        force_new=True 时跳过缓存、强制生成新编号(一行分多批提交时，每批需独立编号，否则同号冲突)。"""
        try:
            user_name = self.get_user_pname()

            # 构建缓存键
            cache_key = f"{user_name}_{method_name}" if method_name else user_name

            # 如果已经为该用户和方法生成过编号，直接返回(force_new 时跳过缓存)
            if not force_new and cache_key in self.experiment_code_cache:
                return self.experiment_code_cache[cache_key]

            # 缓存用户前缀
            if user_name not in self.user_prefix_cache:
                user_prefix = user_name[:3].lower() if user_name else "unk"
                self.user_prefix_cache[user_name] = user_prefix
            else:
                user_prefix = self.user_prefix_cache[user_name]

            # 时间戳为秒级，同秒多批(如 XRF 多方法各一批)会撞号→服务端按号归并/回读回退撞号。
            # 保证全局唯一：与已生成编号同秒则逐秒后移，直到不重复(格式不变：前缀+14位数字)
            ts = datetime.now()
            while f"{user_prefix}{ts.strftime('%Y%m%d%H%M%S')}" in self.experiment_code_cache.values():
                ts += timedelta(seconds=1)
            experiment_code = f"{user_prefix}{ts.strftime('%Y%m%d%H%M%S')}"

            # 缓存生成的编号
            self.experiment_code_cache[cache_key] = experiment_code

            return experiment_code

        except Exception as e:
            if log_func:
                log_func(f"生成实验编号异常: {str(e)}")
            return f"exp_{int(time.time())}"

    def get_experiment_code_for_method(self, method_name, log_func=None):
        """获取特定方法对应的实验编号（如果已生成）"""
        try:
            user_name = self.get_user_pname()
            cache_key = f"{user_name}_{method_name}" if method_name else user_name
            code = self.experiment_code_cache.get(cache_key)

            if code and log_func:
                log_func(f"获取已存在的实验编号: {code} (方法: {method_name})")

            return code

        except Exception as e:
            if log_func:
                log_func(f"获取实验编号异常: {str(e)}")
            return None

    def clear_experiment_codes(self, user_name=None, method_name=None, log_func=None):
        """清理实验编号缓存"""
        try:
            if user_name is None:
                self.experiment_code_cache.clear()
                self.user_prefix_cache.clear()
            elif method_name is None:
                # 清理该用户的所有缓存
                keys_to_remove = [k for k in self.experiment_code_cache.keys() if k.startswith(f"{user_name}_")]
                for key in keys_to_remove:
                    del self.experiment_code_cache[key]
                if user_name in self.user_prefix_cache:
                    del self.user_prefix_cache[user_name]
            else:
                # 清理特定用户和方法的缓存
                cache_key = f"{user_name}_{method_name}"
                if cache_key in self.experiment_code_cache:
                    del self.experiment_code_cache[cache_key]

            if log_func:
                if user_name and method_name:
                    log_func(f"已清理实验编号缓存: {user_name} - {method_name}")
                elif user_name:
                    log_func(f"已清理用户 {user_name} 的所有实验编号缓存")
                else:
                    log_func("已清理所有实验编号缓存")

        except Exception as e:
            if log_func:
                log_func(f"清理实验编号缓存异常: {str(e)}")

    def get_user_pid(self):
        """动态获取当前用户的pid"""
        if not self.login_system.current_user:
            return '377'  # 默认值

        user_info = self.login_system.users.get(self.login_system.current_user, {})
        return user_info.get('pid', '377')

    def get_user_login_id(self):
        """动态获取当前用户的loginId"""
        return self.get_user_pid()

    def get_user_pname(self):
        """获取当前用户名"""
        return self.login_system.current_user or ""

    def load_sub_method_config(self):
        """加载子方法切换配置"""
        try:
            # 假设Excel文件在当前目录
            excel_file = os.path.join(paths.data_dir(), "原始记录登记.xlsx")
            if os.path.exists(excel_file):
                # 使用openpyxl读取Excel
                wb = load_workbook(excel_file)
                if '子方法切换' in wb.sheetnames:
                    sheet = wb['子方法切换']
                    # 从第二行开始读取（跳过标题行）
                    for row in sheet.iter_rows(min_row=2, values_only=True):
                        if row[0] and row[1]:  # 确保主方法ID和子方法ID都不为空
                            main_method_id = str(row[0])
                            sub_method_id = str(row[1])
                            description = row[2] if len(row) > 2 else ""
                            self.sub_method_map[main_method_id] = {
                                'sub_method_id': sub_method_id,
                                'description': description
                            }
        except Exception as e:
            pass

    def get_solution_configure_id_dynamic(self, configure_order, log_func=None):
        """动态查询配置序号对应的配置ID - 严格检查审核状态"""
        try:
            params = {
                "_search": "false",
                "nd": int(time.time() * 1000),
                "pageSize": "100",
                "pageNo": "1",
                "sidx": "",
                "sord": "asc",
                "state": "0",
                "keyword": "",
                "type": "SOLUTION_TYPE_D",
                "pid": self.get_user_pid(),
                "pname": self.get_user_pname(),
                "loginId": self.get_user_login_id()
            }

            headers = {
                'Accept': 'application/json, text/javascript, */*; q=0.01',
                'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; WOW64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/101.0.4951.64 Safari/537.36',
                'Referer': f"{self.login_system.base_url}/web/detectionResultCheckInCalc.html",
                'X-Requested-With': 'XMLHttpRequest'
            }

            response = self.login_system.session.get(
                f"{self.login_system.base_url}/detectionManager/manager/ocExperiment/getSolutionAdata",
                params=params,
                headers=headers,
                verify=False,
                timeout=30
            )

            if response.status_code == 200:
                result = response.json()
                if result.get('success'):
                    vo_list = result.get('resultData', {}).get('voList', [])

                    # 查找匹配的配置序号
                    for solution in vo_list:
                        solution_order = solution.get('configureOrder')
                        if solution_order == configure_order:
                            configure_id = solution.get('id')
                            audit_user_name = solution.get('auditUserName')

                            # 严格的审核状态检查
                            is_audited = audit_user_name is not None and audit_user_name != ""

                            # 如果未审核，直接返回错误
                            if not is_audited:
                                return None, "未审核"

                            if configure_id:
                                return str(configure_id), None

                    return None, "未找到配置序号"
                else:
                    error_msg = result.get('errorCtx', {}).get('errorMsg', '未知错误')
                    return None, error_msg
            else:
                return None, f"HTTP {response.status_code}"

        except Exception as e:
            if log_func:
                log_func(f"动态查询配置序号异常: {str(e)}")
            return None, str(e)


    def get_solution_configure_id_history(self, configure_order, sample_time, log_func=None):
        """查历史标液(dtSolutionConfigure/getSolutionAdata, status=1&configStatus=1)。
        LIMS 前端不支持按 D-XXXX 关键字搜索 → 逐页拉取(每页100, 上限20页)后客户端匹配 configureOrder。
        业务规则：称样时间(分析开始时间) ≤ validityDate 才可选用于录入。
        返回 (configure_id, error_msg)。"""
        try:
            base_params = {
                "_search": "false",
                "pageSize": "100", "sidx": "", "sord": "asc",
                "solutionName": "", "solutionCode": "", "customType": "",
                "configStatus": "1", "controlledNo": "", "storageLocation": "",
                "configureStartDate": "", "configureEndDate": "",
                "configureUserName": "", "receiveUserName": "", "auditStatus": "",
                "status": "1", "type": "SOLUTION_TYPE_D",
                "keyword": "",
                "pid": self.get_user_pid(), "pname": self.get_user_pname(),
                "loginId": self.get_user_login_id(),
            }
            headers = {
                'Accept': 'application/json, text/javascript, */*; q=0.01',
                'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; WOW64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/101.0.4951.64 Safari/537.36',
                'Referer': f"{self.login_system.base_url}/web/detectionResultCheckInCalc.html",
                'X-Requested-With': 'XMLHttpRequest',
            }
            page_size = 100
            total_seen = 0
            # ponytail: 上限20页(2000条)，标液历史超此规模再调大；LIMS 不支持 D-XXXX 搜索故逐页客户端匹配
            for page_no in range(1, 21):
                params = dict(base_params, nd=int(time.time() * 1000), pageNo=str(page_no))
                response = self.login_system.session.get(
                    f"{self.login_system.base_url}/detectionManager/manager/dtSolutionConfigure/getSolutionAdata",
                    params=params, headers=headers, verify=False, timeout=30)
                if response.status_code != 200:
                    return None, f"HTTP {response.status_code}"
                result = response.json()
                if not result.get('success'):
                    return None, (result.get('errorCtx') or {}).get('errorMsg', '未知错误')
                vo_list = result.get('resultData', {}).get('voList', [])
                total_seen += len(vo_list)
                vo = next((s for s in vo_list if s.get('configureOrder') == configure_order), None)
                if vo:
                    if not vo.get('auditUserName'):
                        return None, "未审核"
                    # 有效期校验：称样时间 ≤ validityDate。ISO 日期字符串字典序 = 时间序
                    vd = (vo.get('validityDate') or '').strip()
                    if vd and sample_time is not None:
                        stime = sample_time.strftime("%Y-%m-%d") if hasattr(sample_time, 'strftime') else str(sample_time)[:10]
                        if stime > vd:
                            return None, f"已过期(有效期{vd}，称样时间{stime})"
                    return str(vo.get('id')), None
                if len(vo_list) < page_size:  # 最后一页
                    break
            return None, f"未找到配置序号(历史标液翻页{total_seen}条)"
        except Exception as e:
            if log_func:
                log_func(f"查询历史标液异常: {str(e)}")
            return None, str(e)

    def get_solution_configure_id(self, configure_order, log_func=None, sample_time=None):
        """根据配置序号查询配置ID。先查当前可用标液(ocExperiment)；找不到再查历史标液
        (dtSolutionConfigure)，并按「称样时间 ≤ 有效期」校验。sample_time 用于历史有效期校验。"""
        try:
            configure_id, error_msg = self.get_solution_configure_id_dynamic(configure_order, log_func)

            # 未审核立即返回（当前/历史都不得用未审核标液）
            if error_msg == "未审核":
                if log_func:
                    log_func(f"配置序号 {configure_order} 未审核")
                return None, error_msg

            if configure_id:
                return configure_id, None

            # 当前可用列表没有 → 查历史标液(按有效期校验)
            hid, herr = self.get_solution_configure_id_history(configure_order, sample_time, log_func)
            if hid:
                return hid, None
            return None, herr or error_msg or "未找到配置序号"

        except Exception as e:
            if log_func:
                log_func(f"查询配置序号异常: {str(e)}")
            return None, str(e)

    def submit_solution_with_experiment(self, sample_project_ids, configure_order, experiment_codes, log_func=None):
        """提交标准溶液配置 - 严格审核检查，未审核时完全阻止提交"""
        try:
            # 确保sample_project_ids是逗号分隔的字符串
            if isinstance(sample_project_ids, list):
                sample_project_ids_str = ",".join(str(pid) for pid in sample_project_ids)
            else:
                sample_project_ids_str = str(sample_project_ids)

            # 动态查询配置ID并严格检查审核状态（支持逗号分隔多个标液，每个生成一条 configureJsonList）
            orders = [o.strip() for o in str(configure_order or "").replace("，", ",").split(",") if o.strip()]
            configure_entries = []
            for order in orders:
                configure_id, error_msg = self.get_solution_configure_id(order, log_func)
                # 严格的审核状态检查 - 只有明确返回配置ID才允许提交
                if not configure_id:
                    return False, f"标准溶液 {order} {error_msg}，请先审核后再提交"
                configure_entries.append({
                    "configureIds": str(configure_id),
                    "configureType": "SOLUTION_TYPE_D",
                    "status": 0
                })

            # 获取当前用户信息
            pid = self.get_user_pid()
            pname = self.get_user_pname()
            login_id = self.get_user_login_id()

            # 获取实验编号
            first_project_id = sample_project_ids[0] if isinstance(sample_project_ids, list) else \
            sample_project_ids_str.split(',')[0]
            experiment_code = experiment_codes.get(first_project_id, "")

            if not experiment_code:
                return False, "无法获取实验编号"

            # 获取实验ID
            experiment_id = None
            try:
                config = self.get_experiment_config(
                    first_project_id,
                    "",
                    "",
                    "",
                    log_func
                )
                if config and 'id' in config:
                    experiment_id = config['id']
            except Exception:
                pass

            if not experiment_id:
                experiment_id = 0

            # 构建请求数据
            request_data = {
                "id": int(experiment_id) if experiment_id and str(experiment_id).isdigit() else 0,
                "configureJsonList": json.dumps(configure_entries),
                "pid": pid,
                "pname": pname,
                "loginId": login_id
            }

            # 构建Referer头
            referer_url = f"{self.login_system.base_url}/web/detectionResultCheckInCalc.html?ids=&decideProjectOrgIds=23&type=0&recordNumber=null&checkInStatus=CHECK_IN_STATUS_NO&verifyStatus=&auditStatus=&resultCheckInIds={sample_project_ids_str.split(',')[0] if sample_project_ids_str else ''}&souce=checkIn&pid={pid}&pname={pname}&loginId={login_id}"

            headers = {
                'Content-Type': 'application/json;charset=UTF-8',
                'Accept': 'application/json, text/javascript, */*; q=0.01',
                'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; WOW64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/101.0.4951.64 Safari/537.36',
                'Origin': self.login_system.base_url,
                'Referer': referer_url,
                'Accept-Encoding': 'gzip, deflate',
                'Accept-Language': 'zh-CN,zh;q=0.9'
            }

            # 尝试不同的API端点
            endpoints = [
                "/detectionManager/manager/ocExperiment/saveSolution",
                "/detectionManager/manager/ocExperiment/saveSolutionL"
            ]

            for endpoint in endpoints:
                try:
                    response = self.login_system.session.post(
                        f"{self.login_system.base_url}{endpoint}",
                        json=request_data,
                        headers=headers,
                        verify=False,
                        timeout=30
                    )

                    if response.status_code == 200:
                        try:
                            result = response.json()
                            if result.get('success'):
                                return True, ""
                            else:
                                error_msg = result.get('errorCtx', {}).get('errorMsg', '未知错误')
                                continue
                        except json.JSONDecodeError:
                            continue
                    else:
                        continue

                except Exception:
                    continue

            return False, "所有标准溶液提交端点都失败"

        except Exception as e:
            if log_func:
                log_func(f"提交标准溶液异常: {str(e)}")
            return False, f"提交标准溶液异常: {str(e)}"

    def get_all_solution_types(self, log_func=None):
        """获取所有溶液类型配置 - 处理字典格式，返回键和值"""
        try:
            # 如果已缓存，直接返回
            if self.cached_solution_types:
                return self.cached_solution_types

            params = {
                "pid": self.get_user_pid(),
                "pname": self.get_user_pname(),
                "loginId": self.get_user_login_id(),
                "_": str(int(time.time() * 1000))
            }

            if log_func:
                log_func(f"请求溶液类型URL: /detectionManager/manager/dtSolutionConfigure/getAllSolutionType")
                log_func(f"请求参数: {params}")

            response = self.login_system.session.get(
                f"{self.login_system.base_url}/detectionManager/manager/dtSolutionConfigure/getAllSolutionType",
                params=params,
                headers={
                    'Accept': 'application/json, text/javascript, */*; q=0.01',
                    'Referer': f'{self.login_system.base_url}/web/detectionResultCheckInCalc.html',
                    'X-Requested-With': 'XMLHttpRequest'
                },
                verify=False,
                timeout=30
            )

            if log_func:
                log_func(f"溶液类型响应状态: {response.status_code}")
                log_func(f"溶液类型响应文本: {response.text}")

            if response.status_code == 200:
                result = response.json()
                if result.get('success'):
                    solution_types = result.get('resultData', {})

                    # 处理字典格式
                    if isinstance(solution_types, dict):
                        # 将字典的键和值都作为选项
                        type_options = []
                        for key, value in solution_types.items():
                            type_options.append(key)  # 添加键
                            type_options.append(value)  # 添加值
                        self.cached_solution_types = type_options
                        return type_options
                    elif isinstance(solution_types, list):
                        # 如果是列表，直接返回
                        self.cached_solution_types = solution_types
                        return solution_types
                    else:
                        if log_func:
                            log_func(f"未知的溶液类型数据格式: {type(solution_types)}")
                        return []
                else:
                    error_msg = result.get('errorCtx', {}).get('errorMsg', '未知错误')
                    if log_func:
                        log_func(f"获取溶液类型配置失败: {error_msg}")
                    return []
            else:
                if log_func:
                    log_func(f"获取溶液类型配置失败: HTTP {response.status_code}")
                return []

        except Exception as e:
            if log_func:
                log_func(f"获取溶液类型配置异常: {str(e)}")
            return []

    def ensure_sub_method_and_update(self, sample_project_ids, method_name, method_id, project_names, log_func=None):
        """统一处理子方法切换和更新 - 返回实际使用的方法ID和标准号"""
        try:
            actual_method_id = method_id
            actual_method_name = method_name

            # 检查子方法切换
            if method_id and str(method_id) in self.sub_method_map:
                sub_method_info = self.sub_method_map[str(method_id)]
                actual_method_id = sub_method_info['sub_method_id']
                description = sub_method_info['description']

                # 更新检测方法
                update_success = self.update_method(
                    sample_project_ids,
                    actual_method_id,
                    project_names,
                    log_func
                )

                if update_success:
                    if log_func:
                        log_func(f"已切换到子方法: {description}")
                else:
                    if log_func:
                        log_func(f"警告: 检测方法更新失败，可能无法正确切换子方法")

                # 不覆写 actual_method_name：saveOcExperiment 的 detectionMethod.standardNo
                # 应为父方法标准号（与浏览器行为一致），子方法已通过 update_method 激活
                actual_method_name = method_name

            return actual_method_id, actual_method_name

        except Exception as e:
            if log_func:
                log_func(f"子方法切换异常: {str(e)}")
            return method_id, method_name  # 发生异常时返回原始值

    def get_method_standard_no_by_id(self, method_id, log_func=None):
        """根据方法ID获取方法标准号 - 修复异常处理"""
        try:
            # 如果已经缓存，直接返回
            if str(method_id) in self.method_id_to_standard_no:
                return self.method_id_to_standard_no[str(method_id)]

            # 尝试不同的API端点
            endpoints = [
                "/detectionManager/manager/detectionMethod/getObj",
                "/detectionManager/manager/detectionMethod/pageObj"
            ]

            for endpoint in endpoints:
                try:
                    params = {
                        "id": method_id,
                        "pid": self.get_user_pid(),
                        "pname": self.get_user_pname(),
                        "loginId": self.get_user_login_id(),
                        "_": str(int(time.time() * 1000))
                    }

                    # 对于pageObj端点，需要不同的参数
                    if "pageObj" in endpoint:
                        params = {
                            "pageSize": "100",
                            "pageNo": "1",
                            "id": method_id,
                            "pid": self.get_user_pid(),
                            "pname": self.get_user_pname(),
                            "loginId": self.get_user_login_id(),
                            "_": str(int(time.time() * 1000))
                        }

                    response = self.login_system.session.get(
                        f"{self.login_system.base_url}{endpoint}",
                        params=params,
                        headers={
                            'Accept': 'application/json, text/javascript, */*; q=0.01',
                            'Referer': f'{self.login_system.base_url}/web/detectionResultCheckInCalc.html',
                            'X-Requested-With': 'XMLHttpRequest'
                        },
                        verify=False,
                        timeout=30
                    )

                    if response.status_code == 200:
                        result = response.json()
                        if result.get('success'):
                            # 处理不同的响应格式
                            method_data = None
                            if 'resultData' in result:
                                method_data = result.get('resultData', {})

                                # 处理数组响应
                                if isinstance(method_data, list) and method_data:
                                    method_data = method_data[0]

                                # 处理分页响应
                                elif isinstance(method_data, dict) and 'voList' in method_data:
                                    vo_list = method_data.get('voList', [])
                                    if vo_list:
                                        method_data = vo_list[0]

                            if method_data:
                                standard_no = method_data.get('standardNo') or method_data.get('standard_no')
                                if standard_no:
                                    self.method_id_to_standard_no[str(method_id)] = standard_no
                                    return standard_no

                except Exception as e:
                    continue

            return None

        except Exception as e:
            return None

    def submit_experiment_data(self, experiment_data, project_name, log_func=None, require_signature=False):
        """提交实验数据 - 修复版本，确保与前端请求一致

        require_signature=True 时改调 submitOcExperiment（提交签名/推进工作流），
        否则 saveOcExperiment（仅保存）。两接口请求体一致。
        """
        try:
            if not experiment_data:
                return False, None

            # 构建正确的请求头 - 与前端完全一致
            sample_project_ids = experiment_data.get("sampleProjectIds", "")
            first_project_id = sample_project_ids.split(",")[0] if sample_project_ids else ""

            headers = {
                'Content-Type': 'application/json;charset=UTF-8',
                'Accept': 'application/json, text/javascript, */*; q=0.01',
                'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; WOW64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/101.0.4951.64 Safari/537.36',
                'Origin': self.login_system.base_url,
                'Referer': f'{self.login_system.base_url}/web/detectionResultCheckInCalc.html?ids=&decideProjectOrgIds=23&type=0&recordNumber=null&checkInStatus=CHECK_IN_STATUS_NO&verifyStatus=&auditStatus=&resultCheckInIds={first_project_id}&souce=checkIn&pid={experiment_data.get("pid", self.get_user_pid())}&pname={experiment_data.get("pname", self.get_user_pname())}&loginId={experiment_data.get("loginId", self.get_user_login_id())}',
                'Accept-Encoding': 'gzip, deflate',
                'Accept-Language': 'zh-CN,zh;q=0.9'
            }

            endpoint = "submitOcExperiment" if require_signature else "saveOcExperiment"
            # 大批(百余条) saveOcExperiment/submitOcExperiment 服务端均可能>30s(暂存重跑实测偶>30s 读超时)；
            # 读超时统一给 90s 余量，连接超时维持 10s。
            _timeout = (10, 90)
            response = self.login_system.session.post(
                f"{self.login_system.base_url}/detectionManager/manager/ocExperiment/{endpoint}",
                json=experiment_data,
                headers=headers,
                verify=False,
                timeout=_timeout
            )

            if response.status_code == 200:
                try:
                    result = response.json()
                    if result.get('success'):
                        return True, result.get('resultData')
                    else:
                        error_msg = result.get('errorCtx', {}).get('errorMsg', '未知错误')
                        if log_func:
                            log_func(f"提交实验数据失败: {error_msg}, 项目: {project_name}")
                        return False, f"服务端拒绝: {error_msg}"
                except json.JSONDecodeError:
                    if log_func:
                        log_func(f"提交实验数据失败: 响应不是有效的JSON格式")
                    return False, "响应不是有效的JSON格式"
            else:
                if log_func:
                    _body = (response.text or "")[:500]
                    log_func(f"提交实验数据失败: HTTP {response.status_code}, 响应: {_body}, 项目: {project_name}")
                return False, f"HTTP {response.status_code}"

        except Exception as e:
            if log_func:
                log_func(f"提交实验数据异常: {str(e)}, 项目: {project_name}")
            return False, f"异常: {type(e).__name__}: {str(e)[:120]}"

    def extract_experiment_code(self, result_data, local_code):
        """从提交响应里递归找服务端真实实验编号（与 local_code 同前缀、同长度、后段为数字）；
        找不到回退 local_code。实验编号由服务端生成，本地时间戳拼的会被忽略。"""
        if not local_code or len(local_code) <= 14:
            return local_code
        prefix = local_code[:-14]
        plen = len(local_code)

        def is_code(v):
            return (isinstance(v, str) and len(v) == plen
                    and v.startswith(prefix) and v[len(prefix):].isdigit())

        found = []

        def walk(v):
            if isinstance(v, dict):
                for x in v.values():
                    walk(x)
            elif isinstance(v, list):
                for x in v:
                    walk(x)
            elif is_code(v):
                found.append(v)

        walk(result_data)
        for c in found:
            if c != local_code:
                return c
        return found[0] if found else local_code

    def upload_spectrum_file(self, file_path, result_checkin_ids="", log_func=None):
        """上传谱图文件到 LIMS，返回 (success, inner_resultData_or_errmsg)

        inner 含 id / orgName / name / url；fileName 用 orgName（原始名）。
        body 为 JSON，含 files(data_url)+fileName+type+size+name+pid+pname+loginId
        （抓包 session.txt 确认；缺 fileName 等字段即报"缺少参数"）。URL 无 query；
        result_checkin_ids 仅写入 Referer（服务端实际从 body 取参，其值不影响结果）。
        """
        try:
            if not file_path or not os.path.isfile(file_path):
                return False, f"文件不存在: {file_path}"

            file_name = os.path.basename(file_path)
            file_size = os.path.getsize(file_path)
            with open(file_path, 'rb') as f:
                b64 = base64.b64encode(f.read()).decode('ascii')
            mime, _ = mimetypes.guess_type(file_path)
            if not mime:
                mime = 'application/octet-stream'
            data_url = f"data:{mime};base64,{b64}"

            pid = self.get_user_pid()
            pname = self.get_user_pname()
            login_id = self.get_user_login_id()
            user_info = self.login_system.users.get(self.login_system.current_user, {})
            real_name = user_info.get('display_name') or pname
            # 抓包(session.txt)确认：URL 无 query；body 为含 8 字段的 JSON（files+fileName+type+size+name+pid+pname+loginId）。
            headers = {
                'Content-Type': 'application/json;charset=UTF-8',
                'Accept': 'application/json, text/javascript, */*; q=0.01',
                'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; WOW64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/101.0.4951.64 Safari/537.36',
                'Origin': self.login_system.base_url,
                'Referer': f'{self.login_system.base_url}/web/detectionResultCheckInCalc.html?ids=&decideProjectOrgIds=23&type=0&recordNumber=null&checkInStatus=CHECK_IN_STATUS_NO&verifyStatus=&auditStatus=&resultCheckInIds={result_checkin_ids}&souce=checkIn&pid={pid}&pname={pname}&loginId={login_id}',
                'Accept-Encoding': 'gzip, deflate',
                'Accept-Language': 'zh-CN,zh;q=0.9',
                # 抓包无此头；置 None 剥离 session 默认头（login 时注入）以贴合抓包
                'X-Requested-With': None
            }
            if log_func:
                log_func(f"uploadFile fileName={file_name}, size={file_size}, name={real_name}, resultCheckInIds={result_checkin_ids!r}")

            body = json.dumps({
                "files": data_url,
                "fileName": file_name,
                "type": "base64",
                "size": file_size,
                "name": real_name,
                "pid": int(pid),
                "pname": pname,
                "loginId": int(login_id),
            }, separators=(',', ':')).encode('utf-8')

            response = self.login_system.session.post(
                f"{self.login_system.base_url}/detectionManager/manager/samplePhoto/uploadFile",
                data=body,
                headers=headers,
                verify=False,
                timeout=60
            )

            if response.status_code == 200:
                try:
                    result = response.json()
                    if result.get('success'):
                        inner = result.get('resultData', {}).get('resultData', {})
                        if inner:
                            return True, inner
                        return True, result.get('resultData')
                    error_msg = result.get('errorCtx', {}).get('errorMsg', '未知错误')
                    if log_func:
                        log_func(f"上传谱图失败: {error_msg}, 文件: {os.path.basename(file_path)}")
                    return False, error_msg
                except json.JSONDecodeError:
                    if log_func:
                        log_func("上传谱图失败: 响应不是有效的JSON格式")
                    return False, "响应不是有效的JSON格式"
            else:
                body = response.text[:300] if response.text else ""
                if log_func:
                    log_func(f"上传谱图失败: HTTP {response.status_code}, 文件: {os.path.basename(file_path)}, 响应: {body}")
                return False, f"HTTP {response.status_code}: {body}"

        except Exception as e:
            if log_func:
                log_func(f"上传谱图异常: {str(e)}, 文件: {file_path}")
            return False, str(e)

    def clear_experiment_cache(self, sample_project_ids, log_func=None):
        """清除实验暂存数据 - 完整修复，确保传递所有项目ID"""
        try:
            # 确保sample_project_ids是逗号分隔的字符串
            if isinstance(sample_project_ids, list):
                sample_project_ids_str = ",".join(str(pid) for pid in sample_project_ids)
            else:
                sample_project_ids_str = str(sample_project_ids)

            # 使用抓包获得的正确API端点
            endpoint = "/detectionManager/manager/ocExperiment/cancleOcExperiment"

            # 使用表单数据格式，与抓包请求一致
            request_data = {
                "sampleProjectIds": sample_project_ids_str,
                "pid": self.get_user_pid(),
                "pname": self.get_user_pname(),
                "loginId": self.get_user_login_id()
            }

            # 构建正确的Referer头 - 关键修复
            referer_params = "ids=&decideProjectOrgIds=23&type=0&recordNumber=null&checkInStatus=CHECK_IN_STATUS_NO&verifyStatus=&auditStatus="

            if sample_project_ids_str:
                referer_params += f"&resultCheckInIds={sample_project_ids_str.split(',')[0] if sample_project_ids_str else ''}"

            referer_params += "&souce=checkIn"

            headers = {
                'Accept': 'application/json, text/javascript, */*; q=0.01',
                'Content-Type': 'application/x-www-form-urlencoded; charset=UTF-8',
                'Referer': f'{self.login_system.base_url}/web/detectionResultCheckInCalc.html?{referer_params}',
                'X-Requested-With': 'XMLHttpRequest',
                'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; WOW64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/101.0.4951.64 Safari/537.36'
            }

            response = self.login_system.session.post(
                f"{self.login_system.base_url}{endpoint}",
                data=request_data,
                headers=headers,
                verify=False,
                timeout=30
            )

            if response.status_code == 200:
                result = response.json()
                if result.get('success'):
                    # 删除成功日志
                    return True
                else:
                    error_msg = result.get('errorCtx', {}).get('errorMsg', '未知错误')
                    if log_func:
                        log_func(f"清除实验暂存数据失败: {error_msg}, 项目ID: {sample_project_ids_str}")
                    return False
            else:
                if log_func:
                    log_func(f"清除实验暂存数据失败: HTTP {response.status_code}, 项目ID: {sample_project_ids_str}")
                return False

        except Exception as e:
            if log_func:
                log_func(f"清除实验暂存数据异常: {str(e)}, 项目ID: {sample_project_ids}")
            return False

    def delete_spectrum_by_project_ids(self, project_ids, log_func=None):
        """按项目ID删除已绑定谱图(deleteSpectrumByProjectIds)。
        project_ids 为逗号分隔串或列表；ids=项目ID,pid/pname/loginId=当前登录用户。
        录入前清空旧谱图，避免与新谱图重复。"""
        try:
            if isinstance(project_ids, list):
                ids_str = ",".join(str(p) for p in project_ids if p)
            else:
                ids_str = str(project_ids or "").strip(",")
            if not ids_str:
                return False

            request_data = {
                "ids": ids_str,
                "pid": self.get_user_pid(),
                "pname": self.get_user_pname(),
                "loginId": self.get_user_login_id()
            }

            headers = {
                'Accept': 'application/json, text/javascript, */*; q=0.01',
                'Content-Type': 'application/x-www-form-urlencoded; charset=UTF-8',
                'Referer': f'{self.login_system.base_url}/web/detectionResultCheckInCalc.html?ids=&decideProjectOrgIds=23&type=0&recordNumber=null&checkInStatus=CHECK_IN_STATUS_NO&verifyStatus=&auditStatus=&resultCheckInIds={ids_str}&souce=checkIn&pid={self.get_user_pid()}&pname={self.get_user_pname()}&loginId={self.get_user_login_id()}',
                'X-Requested-With': 'XMLHttpRequest',
                'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; WOW64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/101.0.4951.64 Safari/537.36'
            }

            response = self.login_system.session.post(
                f"{self.login_system.base_url}/detectionManager/manager/ocSpectrum/deleteSpectrumByProjectIds",
                data=request_data,
                headers=headers,
                verify=False,
                timeout=30
            )

            if response.status_code == 200:
                result = response.json()
                if result.get('success'):
                    if log_func:
                        log_func(f"已清空旧谱图(ids={ids_str})")
                    return True
                error_msg = result.get('errorCtx', {}).get('errorMsg', '未知错误')
                if log_func:
                    log_func(f"清空旧谱图失败: {error_msg}, 项目ID: {ids_str}")
                return False
            if log_func:
                log_func(f"清空旧谱图失败: HTTP {response.status_code}, 项目ID: {ids_str}")
            return False

        except Exception as e:
            if log_func:
                log_func(f"清空旧谱图异常: {str(e)}, 项目ID: {project_ids}")
            return False

    def get_units_config(self, log_func=None):
        """获取单位配置 - 使用正确的API地址"""
        try:
            # 如果已缓存，直接返回
            if self.cached_units:
                return self.cached_units

            # 使用正确的API地址和参数
            params = {
                "pageSize": "999",
                "pageNo": "1",
                "reportUnit": "YES",
                "pid": self.get_user_pid(),
                "pname": self.get_user_pname(),
                "loginId": self.get_user_login_id(),
                "_": str(int(time.time() * 1000))
            }

            response = self.login_system.session.get(
                f"{self.login_system.base_url}/detectionManager/manager/meteringUnit/pageObj",
                params=params,
                headers={
                    'Accept': 'application/json, text/javascript, */*; q=0.01',
                    'Referer': f'{self.login_system.base_url}/web/detectionResultCheckInCalc.html',
                    'X-Requested-With': 'XMLHttpRequest'
                },
                verify=False,
                timeout=30
            )

            if response.status_code == 200:
                result = response.json()
                if result.get('success'):
                    units_data = result.get('resultData', {}).get('voList', [])
                    # 缓存单位配置
                    self.cached_units = units_data
                    return units_data
                else:
                    raise Exception("获取单位配置失败，API返回success=false")
            else:
                raise Exception(f"获取单位配置失败: HTTP {response.status_code}")

        except Exception as e:
            raise

    def get_round_methods_config(self, log_func=None):
        """获取修约方法配置"""
        try:
            # 如果已缓存，直接返回
            if self.cached_round_methods:
                return self.cached_round_methods

            # 从系统枚举获取修约方法
            round_methods = {
                "ROUND_METHOD_ENUM_1": "保留{}位有效位数",
                "ROUND_METHOD_ENUM_2": "保留小数点后{}位",
                "ROUND_METHOD_ENUM_3": "保留整数以{}临界点",
                "ROUND_METHOD_ENUM_4": "科学计数法表示"
            }

            round_keep_methods = {
                "ROUND_KEEP_ENUM_1": "四舍五入",
                "ROUND_KEEP_ENUM_2": "四舍六入五留双",
                "ROUND_KEEP_ENUM_3": ">0 直接进位",
                "ROUND_KEEP_ENUM_4": "直接截取"
            }

            config = {
                "round_methods": round_methods,
                "round_keep_methods": round_keep_methods
            }

            # 缓存配置
            self.cached_round_methods = config
            return config

        except Exception as e:
            raise

    def get_calc_methods_config(self, log_func=None):
        """获取计算方法配置"""
        try:
            # 如果已缓存，直接返回
            if self.cached_calc_methods:
                return self.cached_calc_methods

            # 从系统枚举获取计算方法
            calc_methods = {
                "CALC_METHOD_ENUM_AVG": "按平均值方式",
                "CALC_METHOD_ENUM_MAX": "按最大值方式",
                "CALC_METHOD_ENUM_GATHER": "按汇总方式",
                "CALC_METHOD_ENUM_SUM": "按求和方式",
                "CALC_METHOD_ENUM_MEDIAN": "按中位值方式",
                "CALC_METHOD_ENUM_MPN": "MPN检索表",
                "CALC_METHOD_MNUM_JOIN": "按+获取表头",
                "CALC_METHOD_CHOICE_CONTENT": "按勾选内容"
            }

            # 缓存配置
            self.cached_calc_methods = calc_methods
            return calc_methods

        except Exception as e:
            raise

    def get_method_specific_dynamic_columns(self, sample_project_ids, method_standard_no, method_id=None,
                                            log_func=None):
        """获取特定检测方法的动态列配置 - 修复子方法切换问题"""
        if not self.login_system.current_user:
            raise Exception("用户未登录")

        try:
            # 保存原始参数用于回退
            original_method_id = method_id
            original_standard_no = method_standard_no

            # 如果有方法ID，检查是否需要子方法切换
            actual_method_id = method_id
            if method_id:
                actual_method_id = self.check_and_switch_sub_method(method_id, log_func)

            # 使用 Fiddler 抓包获得的完整参数
            request_data = {
                "sampleProjectIds": sample_project_ids,
                "experimentCode": "",
                "detectionMethod.standardNo": method_standard_no,
                "detectionMethod.id": actual_method_id if actual_method_id else "",
                "detectionMethod.subMethodName": "",
                "pid": self.get_user_pid(),
                "pname": self.get_user_pname(),
                "loginId": self.get_user_login_id(),
                "isCalculatedValue": "1",
                "sampleEvaluate": "0",
                "reportUnit": "1",
                "isOptionalTable": "0"
            }

            # 第一次尝试：使用当前参数获取动态列
            response = self.login_system.session.post(
                f"{self.login_system.base_url}/detectionManager/manager/ocExperiment/getDynamicColumns",
                data=request_data,
                headers={
                    'Accept': 'application/json, text/javascript, */*; q=0.01',
                    'Content-Type': 'application/x-www-form-urlencoded; charset=UTF-8',
                    'Referer': f'{self.login_system.base_url}/web/detectionResultCheckInCalc.html',
                    'X-Requested-With': 'XMLHttpRequest',
                    'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; WOW64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/101.0.4951.64 Safari/537.36'
                },
                verify=False,
                timeout=30
            )

            # 处理响应
            if response.status_code == 200:
                result = response.json()
                if result.get('success'):
                    dynamic_columns = result.get('resultData', [])
                    return dynamic_columns
                else:
                    error_msg = result.get('errorCtx', {}).get('errorMsg', '未知错误')
                    # 如果是下拉菜单未设置选项的错误，尝试回退策略
                    if "动态列为下拉菜单，但未设置下拉菜单的选项" in error_msg:
                        return self.fallback_dynamic_columns(
                            sample_project_ids, method_standard_no, method_id,
                            original_method_id, original_standard_no,
                            request_data, log_func
                        )
                    else:
                        # 其他错误返回空数组
                        return []
            else:
                # HTTP错误时也尝试回退
                return self.fallback_dynamic_columns(
                    sample_project_ids, method_standard_no, method_id,
                    original_method_id, original_standard_no,
                    request_data, log_func
                )

        except Exception as e:
            # 异常情况下返回空数组，避免整个流程中断
            return []

    def fallback_dynamic_columns(self, sample_project_ids, method_standard_no, method_id,
                                 original_method_id, original_standard_no, request_data, log_func):
        """回退策略：尝试不同的方法获取动态列"""

        # 策略1: 如果发生了子方法切换，回退到主方法
        if method_id and original_method_id and method_id != original_method_id:
            request_data["detectionMethod.id"] = original_method_id
            response = self.retry_dynamic_columns_request(request_data, log_func)
            if response is not None:
                return response

        # 策略2: 尝试不传递方法ID
        request_data["detectionMethod.id"] = ""
        response = self.retry_dynamic_columns_request(request_data, log_func)
        if response is not None:
            return response

        # 策略3: 尝试使用原始标准号和方法ID
        request_data["detectionMethod.standardNo"] = original_standard_no
        request_data["detectionMethod.id"] = original_method_id if original_method_id else ""
        response = self.retry_dynamic_columns_request(request_data, log_func)
        if response is not None:
            return response

        # 所有策略都失败，返回空数组
        return []

    def retry_dynamic_columns_request(self, request_data, log_func):
        """重试动态列请求"""
        try:
            response = self.login_system.session.post(
                f"{self.login_system.base_url}/detectionManager/manager/ocExperiment/getDynamicColumns",
                data=request_data,
                headers={
                    'Accept': 'application/json, text/javascript, */*; q=0.01',
                    'Content-Type': 'application/x-www-form-urlencoded; charset=UTF-8',
                    'Referer': f'{self.login_system.base_url}/web/detectionResultCheckInCalc.html',
                    'X-Requested-With': 'XMLHttpRequest',
                    'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; WOW64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/101.0.4951.64 Safari/537.36'
                },
                verify=False,
                timeout=30
            )

            if response.status_code == 200:
                result = response.json()
                if result.get('success'):
                    dynamic_columns = result.get('resultData', [])
                    return dynamic_columns

            return None

        except Exception as e:
            return None

    def calc_report_values(self, calc_rows, meta, log_func=None):
        """调用 ocMethodTitleSettings/calcTheValue 计算报告值，返回 {row_id: {calculatedValue, reportValue, other}}。
        calc_rows: [{columnValues, id, projectId, sampleCode, calcUnit, reportUnit, serialNum}, ...]（id 与分析记录 id 一致）
        meta: {roundMethod, roundMethodLevelJson, resultRoundMethod, resultRoundMethodLevelJson,
               calcMethod, accuracy, methodSettingId, accuracyRadixPoint, curve}
        响应 resultData 是 JSON 字符串需二次解析；失败/异常返回 {}，调用方回退原 None 行为不阻断提交。
        注：单位换算(unitConversion/conversionBatch)按需，数值报告值单位不符时再补。"""
        try:
            pid = self.get_user_pid()
            try:
                pid = int(pid)
            except (ValueError, TypeError):
                pid = 377
            curve = meta.get('curve')
            if not curve:
                curve = [{"id": "", "equation": "", "correlationCoefficient": "", "slope": "", "intercept": ""}]
            elif isinstance(curve, dict):
                curve = [curve]
            body = {
                "roundMethod": meta.get('roundMethod'),
                "roundMethodLevelJson": meta.get('roundMethodLevelJson'),
                "resultRoundMethod": meta.get('resultRoundMethod'),
                "resultRoundMethodLevelJson": meta.get('resultRoundMethodLevelJson'),
                "calcMethod": meta.get('calcMethod'),
                "accuracy": meta.get('accuracy'),
                "methodSettingId": meta.get('methodSettingId'),
                "accuracyRadixPoint": meta.get('accuracyRadixPoint'),
                "curve": curve,
                "rows": calc_rows,
                "pid": pid,
                "pname": self.get_user_pname(),
                "loginId": pid,
            }
            response = self.login_system.session.post(
                f"{self.login_system.base_url}/detectionManager/manager/ocMethodTitleSettings/calcTheValue",
                json=body,
                headers={
                    'Content-Type': 'application/json;charset=UTF-8',
                    'Accept': 'application/json, text/javascript, */*; q=0.01',
                    'Origin': self.login_system.base_url,
                    'Referer': f'{self.login_system.base_url}/web/detectionResultCheckInCalc.html',
                    'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; WOW64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/101.0.4951.64 Safari/537.36',
                },
                verify=False, timeout=30,
            )
            if response.status_code != 200:
                if log_func:
                    log_func(f"calcTheValue 失败: HTTP {response.status_code}; 响应={response.text[:200]}")
                return {}
            result = response.json()
            if not result.get('success'):
                err = result.get('errorCtx') or {}
                if log_func:
                    log_func(f"calcTheValue 失败(success=false): {err.get('errorMsg', '未知错误')}; 响应={response.text[:200]}")
                return {}
            rd = result.get('resultData')
            if isinstance(rd, str):
                rd = json.loads(rd)  # resultData 是 JSON 字符串，需二次解析
            out = {}
            for row in (rd or {}).get('rows', []):
                cv = row.get('columnValues', {})
                out[row.get('id')] = {
                    'calculatedValue': cv.get('calculatedValue'),
                    'reportValue': cv.get('reportValue'),
                    'other': cv.get('other'),
                }
            if log_func and not out:
                log_func(f"calcTheValue: 响应 0 行 (rd 类型={type(rd).__name__})")
            return out
        except Exception as e:
            if log_func:
                log_func(f"calcTheValue 异常: {str(e)}")
            return {}

    def get_experiment_config(self, sample_project_ids, method_standard_no, result_checkin_ids=None, sample_id=None,
                              log_func=None, method_id=None):
        """获取实验配置 - 确保使用子方法配置"""
        try:
            # 验证参数
            if not sample_project_ids:
                raise Exception("项目ID列表为空，无法获取实验配置")

            # 检查并切换子方法
            actual_method_id = method_id
            if method_id:
                actual_method_id = self.check_and_switch_sub_method(method_id, log_func)

            # 使用完整的项目ID列表
            request_data = {
                "sampleProjectIds": sample_project_ids,
                "calcType": "0",
                # 方法标准号：让服务端定位检测方法，从而展开平行样记录并回填 methodId；
                # 与 getDynamicColumns 保持一致（之前漏传导致 methodId 未解析、平行样不展开）
                "detectionMethod.standardNo": method_standard_no,
                "pid": self.get_user_pid(),
                "pname": self.get_user_pname(),
                "loginId": self.get_user_login_id(),
            }

            # 只有在有值的情况下才添加可选参数
            if result_checkin_ids:
                request_data["resultCheckInIds"] = result_checkin_ids
            if sample_id:
                request_data["sampleId"] = sample_id
            # 添加方法ID参数 - 使用子方法ID
            if actual_method_id:
                request_data["detectionMethod.id"] = actual_method_id

            # 构建Referer头
            referer_params = "ids=&decideProjectOrgIds=23&type=0&recordNumber=null&checkInStatus=CHECK_IN_STATUS_NO&verifyStatus=&auditStatus="

            if result_checkin_ids:
                referer_params += f"&resultCheckInIds={result_checkin_ids}"
            if sample_id:
                referer_params += f"&sampleId={sample_id}"

            referer_params += "&souce=checkIn"

            headers = {
                'Accept': 'application/json, text/javascript, */*; q=0.01',
                'Content-Type': 'application/x-www-form-urlencoded; charset=UTF-8',
                'Referer': f'{self.login_system.base_url}/web/detectionResultCheckInCalc.html?{referer_params}',
                'X-Requested-With': 'XMLHttpRequest',
                'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; WOW64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/101.0.4951.64 Safari/537.36'
            }

            response = self.login_system.session.post(
                f"{self.login_system.base_url}/detectionManager/manager/ocExperiment/getOcExperiment",
                data=request_data,
                headers=headers,
                verify=False,
                timeout=30
            )

            if response.status_code == 200:
                result = response.json()
                if result.get('success'):
                    config_data = result.get('resultData', {})
                    return config_data
                else:
                    if log_func:
                        log_func(f"getOcExperiment success=false: {result}")
                    return {}
            else:
                if log_func:
                    log_func(f"getOcExperiment HTTP {response.status_code}: {response.text[:200]}")
                return {}

        except Exception as e:
            return {}

    def get_counterpart_experiment(self, sample_code, counterpart_project, log_func=None):
        """跨序列共享称样量：定位同报验编号已登记的「对应方法」实验，回读其称样量。
        counterpart_project = 对应项目名(编辑器 weighing_params.counterpart 填的)。
        流程：查该样品已登记项目(CHECK_IN_STATUS_ALREADY) → projectName 精确匹配定位对应方法那一条
        (projectName 唯一；不用标准号超集定位) → get_all_configs 取其实验记录+动态列元数据。
        返回 {"dynamic_columns":[...], "records":[...], "project_name":"..."} 或 None。
        None 含义：对应方法未登记/未找到/取配置失败 → 调用方落新生成。"""
        try:
            cp = _norm_cn(counterpart_project)
            if not cp:
                return None
            # 精确查该报验编号的已登记项目(同主流程精确匹配)
            projects = self.query_samples_by_conditions(
                sample_code=sample_code, exact_match=True,
                check_in_status="CHECK_IN_STATUS_ALREADY", log_func=log_func,
            ) or []
            located = next(
                (p for p in projects if _norm_cn(p.get("projectName") or "") == cp), None
            )
            if not located:
                if log_func:
                    log_func(f"跨序列对应: {sample_code} 的已登记项目中无「{counterpart_project}」")
                return None
            pid = located.get("projectId")
            if not pid:
                return None
            standard_no = located.get("standardNo") or ""  # 从定位结果现读(跟主流程 projects[0].standardNo 一致)
            method_id = located.get("decideProjectMethodId")
            # get_all_configs 协调子方法切换+实验记录+动态列元数据(含 isWeighing)；抛异常时落 None
            cfg = self.get_all_configs(str(pid), standard_no, "", "", log_func, method_id, [])
            experiment = (cfg or {}).get("experiment") or {}
            records = experiment.get("ocAnalysisRecordList") or []
            dynamic_columns = (cfg or {}).get("dynamic_columns") or []
            if not records or not dynamic_columns:
                if log_func:
                    log_func(f"跨序列对应: {sample_code} 的「{counterpart_project}」取到空记录/空列")
                return None
            return {"dynamic_columns": dynamic_columns, "records": records,
                    "project_name": located.get("projectName") or ""}
        except Exception as e:
            if log_func:
                log_func(f"get_counterpart_experiment 异常: {type(e).__name__}: {str(e)[:200]}")
            return None

    def get_oc_compare_show_data(self, sample_code, items, log_func=None):
        """读取"对比展示"数据：给定样品号 + 组分项目名列表(items)，返回各组分已录入实验记录
        (含 projectName/reportValue/calculatedValue/serialNumber/sampleCode …)。
        用于"总和"项目取各组分报告值。响应双层 resultData：resp['resultData']['resultData'] = 记录列表；失败/空返回 []。"""
        try:
            params = {
                "sampleCodes": sample_code,
                "items": json.dumps(items, ensure_ascii=False),
                "pid": self.get_user_pid(),
                "pname": self.get_user_pname(),
                "loginId": self.get_user_login_id(),
                "_": int(time.time() * 1000),
            }
            headers = {
                'Accept': 'application/json, text/javascript, */*; q=0.01',
                'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; WOW64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/101.0.4951.64 Safari/537.36',
                'Referer': f"{self.login_system.base_url}/web/detectionResultCheckInCalc.html",
                'X-Requested-With': 'XMLHttpRequest',
            }
            response = self.login_system.session.get(
                f"{self.login_system.base_url}/detectionManager/manager/ocExperiment/getOcCompareShowData",
                params=params, headers=headers, verify=False, timeout=30,
            )
            if response.status_code == 200:
                result = response.json()
                if result.get('success'):
                    outer = result.get('resultData') or {}
                    return outer.get('resultData') or []
                if log_func:
                    log_func(f"getOcCompareShowData success=false: {result.get('errorCtx')}")
                return []
            if log_func:
                log_func(f"getOcCompareShowData HTTP {response.status_code}: {response.text[:200]}")
            return []
        except Exception as e:
            if log_func:
                log_func(f"getOcCompareShowData 异常: {e}")
            return []

    def query_samples_by_conditions(self, sample_code=None, project_name=None, method_name=None, retest_checked=False,
                                    log_func=None, exact_match=False, days=30, check_in_status="CHECK_IN_STATUS_NO"):
        """通过多个条件查询样品信息 - 支持精确匹配和模糊查询
        days: 受理日期窗口(天)，默认30；方法池查询可按方法文件配置收窄提速，逐样品精确查保持30(系统可查上限)。
        check_in_status: CHECK_IN_STATUS_NO(未登记，默认) / CHECK_IN_STATUS_ALREADY(已登记)。"""
        if not self.login_system.current_user:
            return []

        try:
            today = datetime.now()
            one_month_ago = today - timedelta(days=days)

            # 优化分页参数
            page_size = 500
            page_no = 1
            all_projects = []

            # 循环获取所有页面的数据
            while True:
                # 直接在API参数中使用所有条件
                params = {
                    "_search": "false",
                    "nd": int(time.time() * 1000),
                    "pageSize": page_size,
                    "pageNo": page_no,
                    "acceptStartDate": one_month_ago.strftime("%Y-%m-%d"),
                    "acceptEndDate": today.strftime("%Y-%m-%d"),
                    "checkInStatus": check_in_status,
                    "decideProjectOrgId": "23",
                    "sampleStatus": "one",
                    "pid": self.get_user_pid(),
                    "pname": self.get_user_pname(),
                    "loginId": self.get_user_login_id(),
                }

                # 精确匹配模式：用报验编号(字母+8位)作 keyword——服务端按 detectionNo 匹配，
                # 输入若带小号/平行字母/后缀(如 TS26080731001_报告)需先归一到报验编号 TS26080731
                if exact_match and sample_code:
                    params["keyword"] = _detection_no_of(sample_code)
                else:
                    # 模糊查询模式：使用各个字段分别查询
                    if sample_code:
                        params["keyword"] = sample_code
                    if project_name:
                        params["decideProjectName"] = project_name
                    if method_name:
                        params["decideProjectMethodName"] = method_name

                if retest_checked:
                    params["cancelRetesting"] = "1"

                response = self.login_system.session.get(
                    f"{self.login_system.base_url}/detectionManager/manager/resultCheckIn/pagePCObjAndSample",
                    params=params,
                    headers={
                        'Accept': 'application/json, text/javascript, */*; q=0.01',
                        'Referer': f'{self.login_system.base_url}/web/detectionResultCheckInListMgt.html'
                    },
                    verify=False,
                    timeout=60
                )

                if response.status_code == 200:
                    result = response.json()
                    if result.get('success') and result.get('resultData', {}).get('voList'):
                        vo_list = result['resultData']['voList']
                        current_page_records = len(vo_list)

                        # 处理当前页的记录
                        for sample_data in vo_list:
                            detection_no = sample_data.get('detectionNo', '')
                            small_no = sample_data.get('smallNo', '')
                            current_sample_code = f"{detection_no}{small_no}"

                            # 判断是否为复测项目
                            old_sample_project_id = sample_data.get('oldSampleProjectId')
                            is_retest = old_sample_project_id is not None

                            project_data = {
                                'sampleCode': current_sample_code,
                                'sampleName': sample_data.get('sampleName', ''),
                                'projectName': sample_data.get('decideProjectName', ''),
                                'standardNo': sample_data.get('standardNo', ''),
                                'decideProjectMethodId': sample_data.get('decideProjectMethodId'),
                                'subMethodName': sample_data.get('subMethodName', ''),
                                'projectId': sample_data.get('id'),
                                # 检测项目定义ID（取项目别名 detailByProject?id= 要用；字段名待 Step 0 确认）
                                'detectionProjectId': sample_data.get('detectionProjectId') or sample_data.get('decideProjectId'),
                                'sampleId': sample_data.get('sampleId'),
                                'detectionNo': detection_no,
                                'sampleSmallNo': small_no,
                                'isRetest': is_retest,
                                'oldSampleProjectId': old_sample_project_id,
                                'checkInStatus': sample_data.get('checkInStatus', ''),
                                '_raw': sample_data  # 透传原始(含受理时间等未解析字段)，供 startTime ≥ 受理时间 校验用
                            }
                            all_projects.append(project_data)

                        # 精确匹配模式：找到完全匹配的样品后立即返回
                        # 早退仅当 sample_code 是完整样品号(含小号)时安全：
                        # 报验编号级 sample_code(如 dedup 传报验编号)会命中 smallNo 为空的
                        # 残留记录(sampleCode 恰为报验编号)而误早退，丢弃该报验编号下其余小号。
                        if exact_match and sample_code and sample_code != _detection_no_of(sample_code):
                            matched_projects = [p for p in all_projects if p.get('sampleCode') == sample_code]
                            if matched_projects:
                                return matched_projects

                        # 检查是否还有更多数据
                        if current_page_records < page_size:
                            break
                        else:
                            page_no += 1
                    else:
                        break
                else:
                    break

            return all_projects

        except Exception as e:
            if log_func:
                log_func(f"[API异常] {type(e).__name__}: {str(e)[:200]}")
            return []

    def diagnose_missing_sample(self, sample_code, log_func=None, days=30):
        """录入查询查不到样品时，按 收样→制样→登记 生命周期定位真实原因，返回具体中文理由。
        串行：项目已登记(录入端点 checkInStatus=ALREADY 重查)→未制样(pageObjByMakeStatus)→未收样(pageObj)→兜底。
        days: 受理日期窗口(天)，与主查询共用方法文件 date_window_days，避免窗口不一致致已登记样品误报"不存在"。
        报验编号=字母前缀+8位流水(_detection_no_of 提取)；各端点 keyword 统一用报验编号。
        异常仅记日志不抛、退回兜底文案，保证不阻断序列运行。"""
        base = self.login_system.base_url
        sess = self.login_system.session
        common = {"pid": self.get_user_pid(), "pname": self.get_user_pname(),
                  "loginId": self.get_user_login_id()}
        headers = {'Accept': 'application/json, text/javascript, */*; q=0.01',
                   'Referer': f'{base}/web/detectionResultCheckInListMgt.html'}
        dno = _detection_no_of(sample_code)  # 报验编号=字母+8位(原 [:-3] 遇 _报告 等后缀会误剥)

        def _hit(url, params):
            r = sess.get(url, params=params, headers=headers, verify=False, timeout=60)
            if r.status_code != 200:
                return False
            _vl = (r.json().get('resultData') or {}).get('voList') or []
            return bool(_vl)

        try:
            # 1) 项目已登记：重查录入端点，checkInStatus 用 ALREADY(非 YES——抓包确认已登记列表页用此枚举)
            today = datetime.now()
            window_ago = today - timedelta(days=days)
            if _hit(f"{base}/detectionManager/manager/resultCheckIn/pagePCObjAndSample",
                    {**common, "_search": "false", "nd": int(time.time() * 1000),
                     "pageSize": 30, "pageNo": 1, "sampleStatus": "one", "decideProjectOrgId": "23",
                     "acceptStartDate": window_ago.strftime("%Y-%m-%d"),
                     "acceptEndDate": today.strftime("%Y-%m-%d"),
                     "checkInStatus": "CHECK_IN_STATUS_ALREADY", "keyword": dno}):
                return "已登记"
            # 制样/收样端点公共参数(按抓包：makeSampleMarkNames=A，keyword=报验编号)
            sp = {"_search": "false", "nd": int(time.time() * 1000), "pageSize": 30, "pageNo": 1,
                  "sidx": "", "sord": "asc", "makeSampleMarkNames": "A", "advanced": "",
                  **common, "keyword": dno}
            # 2) 未制样
            if _hit(f"{base}/detectionManager/manager/sample/pageObjByMakeStatus",
                    {**sp, "sampleMakeStatus": "SAMPLE_MAKE_STATUS_NO"}):
                return "未制样"
            # 3) 未收样
            if _hit(f"{base}/detectionManager/manager/sample/pageObj",
                    {**sp, "sampleReceiveStatus": "SAMPLE_RECEIVE_STATUS_NO_INVENTORY_STATUS_ALREADY"}):
                return "未收样"
            # 3.5) 状态枚举兜底：无过滤按报验编号查 pageObj，读 processStatus 精确归类
            # sampleReceiveStatus 枚举组合不稳(如 NO_INVENTORY_STATUS_ALREADY 只命中部分未收样样品)，
            # 其余未收样样品会漏判→误报"不存在"；processStatus 是 LIMS 返回的权威状态文案。
            r = sess.get(f"{base}/detectionManager/manager/sample/pageObj",
                         params=sp, headers=headers, verify=False, timeout=60)
            if r.status_code == 200:
                _vl = (r.json().get('resultData') or {}).get('voList') or []
                if _vl:
                    _ps = str(_vl[0].get('processStatus') or '').strip()
                    if _ps:
                        return _ps
        except Exception as e:
            if log_func:
                log_func(f"[诊断异常] {type(e).__name__}: {str(e)[:200]}")
        return f"报验编号不存在或超过可查期限(>{days}天)"

    def get_all_configs(self, sample_project_ids, method_standard_no, result_checkin_ids=None, sample_id=None,
                        log_func=None, method_id=None, project_names=None):
        """获取所有配置信息 - 集成动态列处理"""
        if not self.login_system.current_user:
            raise Exception("用户未登录")

        # 参数验证
        if not sample_project_ids:
            raise Exception("项目ID列表为空")

        if not method_standard_no:
            raise Exception("检测方法标准号为空")

        try:
            # 统一处理子方法切换
            actual_method_id, actual_method_name = self.ensure_sub_method_and_update(
                sample_project_ids,
                method_standard_no,
                method_id,
                project_names or [],
                log_func
            )

            configs = {}

            # 1. 获取实验配置 - 使用实际的方法ID
            experiment_config = self.get_experiment_config(
                sample_project_ids,
                actual_method_name,
                result_checkin_ids,
                sample_id,
                log_func,
                actual_method_id
            )
            configs['experiment'] = experiment_config

            # 2. 获取设备信息 - 使用实际的方法ID
            if actual_method_id:
                self.current_sample_id = sample_id
                equipment_config = self.get_detection_equipment(actual_method_id, log_func)
                configs['equipment'] = equipment_config
            else:
                configs['equipment'] = self.get_empty_equipment()

            # 3. 获取单位配置
            configs['units'] = self.get_units_config(log_func)

            # 4. 获取修约方法配置
            configs['round_methods'] = self.get_round_methods_config(log_func)

            # 5. 获取计算方法配置
            configs['calc_methods'] = self.get_calc_methods_config(log_func)

            # 6. 获取动态列配置 - 使用实际的方法ID
            dynamic_columns = self.get_method_specific_dynamic_columns(
                sample_project_ids, actual_method_name, actual_method_id, log_func
            )

            # 7. 如果API返回的动态列为空，尝试从实验配置中提取
            if not dynamic_columns and experiment_config:
                dynamic_columns = self.extract_dynamic_columns_from_config(experiment_config, log_func)

            configs['dynamic_columns'] = dynamic_columns

            # 返回实际使用的方法信息
            configs['actual_method_id'] = actual_method_id
            configs['actual_method_name'] = actual_method_name

            return configs

        except Exception as e:
            raise

    def get_detection_equipment(self, detection_method_id, log_func=None):
        """获取检测方法对应的设备信息 - 修复参数格式问题"""
        try:
            # 检查并切换子方法 - 确保返回的是字符串ID
            original_method_id = detection_method_id
            detection_method_id = self.check_and_switch_sub_method(detection_method_id, log_func)

            # 确保detection_method_id是字符串，不是元组
            if isinstance(detection_method_id, tuple):
                detection_method_id = detection_method_id[0]

            # 构建请求参数 - 确保参数是字符串
            params = {
                "detectionMethodId": str(detection_method_id),
                "pid": self.get_user_pid(),
                "pname": self.get_user_pname(),
                "loginId": self.get_user_login_id(),
                "_": str(int(time.time() * 1000))
            }

            # 构建Referer
            sample_id_param = f"&sampleId={self.current_sample_id}" if self.current_sample_id else ""
            referer_url = f"{self.login_system.base_url}/web/detectionResultCheckInCalc.html?ids=&decideProjectOrgIds=23&type=0&recordNumber=null&checkInStatus=CHECK_IN_STATUS_NO&verifyStatus=&auditStatus=&resultCheckInIds={sample_id_param}&souce=checkIn"

            headers = {
                'Accept': 'application/json, text/javascript, */*; q=0.01',
                'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; WOW64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/101.0.4951.64 Safari/537.36',
                'Referer': referer_url,
                'Accept-Encoding': 'gzip, deflate',
                'Accept-Language': 'zh-CN,zh;q=0.9',
                'X-Requested-With': 'XMLHttpRequest'
            }

            response = self.login_system.session.get(
                f"{self.login_system.base_url}/detectionManager/manager/detectionMethod/selectByDetectionMethodId",
                params=params,
                headers=headers,
                verify=False,
                timeout=30
            )

            if response.status_code == 200:
                result = response.json()
                if result.get('success'):
                    equipment_list = result.get('resultData', [])
                    equipment_data = self.process_equipment_list(equipment_list, log_func)
                    return equipment_data
                else:
                    return self.get_empty_equipment()
            else:
                return self.get_empty_equipment()

        except Exception as e:
            return self.get_empty_equipment()

    def process_equipment_list(self, equipment_list, log_func=None):
        """处理设备列表，提取主检设备和称样设备 - 修复设备信息格式，避免橙色背景"""
        main_equipments = []
        weighing_equipments = []
        main_equipment_names_list = []
        main_equipment_ids_list = []
        weighing_equipment_ids_list = []

        # 遍历设备列表，按用途分类，只选择默认设备
        if log_func and equipment_list:
            log_func("[设备] 方法返回 {} 台: {}".format(
                len(equipment_list), "; ".join(
                    "{}({}/isDefault={!r})".format(e.get('name'), e.get('usedCategory'), e.get('isDefault'))
                    for e in equipment_list)))
        for equipment in equipment_list:
            used_category = equipment.get('usedCategory', '')
            # 设备名逐级回落：selectByDetectionMethodId 返回里称样设备 name=None，真名在 equipmentBillName
            name = (equipment.get('name') or equipment.get('equipmentName')
                    or equipment.get('equipmentBillName') or '').strip()
            main_equipment_names_value = equipment.get('mainEquipmentNames', '')
            # 显示名优先 mainEquipmentNames(已含"编号,名称,日期")；缺失时按 GUI(detection_entry_main:1989)
            # 同款用 no/name/checkOutDate 拼，使称样设备也能得到"编号,名称,日期"串
            if main_equipment_names_value:
                display_name = main_equipment_names_value
            else:
                _eno = (equipment.get('no') or '').strip()
                _cod = equipment.get('checkOutDate')
                _edate = _cod[:10] if isinstance(_cod, str) else (str(_cod)[:10] if _cod else '')
                display_name = "{},{},{}".format(_eno, name, _edate) if _eno else name
            is_default = equipment.get('isDefault')
            equipment_id = equipment.get('id')
            equipment_bill_id = equipment.get('equipmentBillId')

            # 优先使用 equipmentBillId，如果不存在则使用 id
            effective_id = equipment_bill_id if equipment_bill_id else equipment_id

            # 只处理默认设备 (isDefault=1；兼容字符串"1"/布尔True，否则字符串"1"会被跳过→默认设备填不上)
            if str(is_default).strip().lower() not in ("1", "true", "是"):
                continue
            if log_func and not display_name:
                log_func("[设备] 默认设备无法解析名称，原始字段: {}".format(
                    {k: v for k, v in equipment.items() if v not in (None, '', [], {})}))

            if used_category == '检测设备':
                equipment_info = {
                    'name': name,
                    'display_name': display_name,
                    'mainEquipmentNames': main_equipment_names_value,
                    'isDefault': is_default,
                    'equipment': equipment,
                    'id': effective_id
                }
                main_equipments.append(equipment_info)

                # 记录显示名称和ID
                main_equipment_names_list.append(display_name)
                if effective_id:
                    main_equipment_ids_list.append(str(effective_id))

            elif used_category == '称样设备':
                equipment_info = {
                    'name': name,
                    'display_name': display_name,
                    'isDefault': is_default,
                    'equipment': equipment,
                    'id': effective_id
                }
                weighing_equipments.append(equipment_info)

                if effective_id:
                    weighing_equipment_ids_list.append(str(effective_id))

        # 选择主检设备 - 只选择默认设备
        if main_equipments:
            # 构建主检设备显示名称列表
            main_equipment_display_names = [eq['display_name'] for eq in main_equipments]

            # 用分号连接所有默认主检设备
            main_equipment = "; ".join(main_equipment_display_names)
            main_equipment_names = "; ".join(main_equipment_names_list)
            main_equipment_ids = ",".join(main_equipment_ids_list)
        else:
            main_equipment = ""
            main_equipment_names = ""
            main_equipment_ids = ""

        # 选择称样设备 - 只选择默认设备
        weighing_equipment_base_name = ""
        weighing_equipment_display_name = ""
        weighing_equipment_id = ""

        if weighing_equipments:
            # 只选择默认称样设备
            selected_weighing = weighing_equipments[0]
            weighing_equipment_name = selected_weighing['name'] or ""
            weighing_equipment_display_name = selected_weighing['display_name'] or ""
            weighing_equipment_id = selected_weighing['id']

            # 提取设备基本名称（去掉编码和日期信息）
            if ',' in weighing_equipment_display_name:
                parts = weighing_equipment_display_name.split(',')
                if len(parts) >= 2:
                    weighing_equipment_base_name = parts[1]
                else:
                    weighing_equipment_base_name = weighing_equipment_name
            else:
                weighing_equipment_base_name = weighing_equipment_name
        else:
            weighing_equipment_display_name = ""
            weighing_equipment_base_name = ""
            weighing_equipment_id = ""

        # 修复：返回与前端操作一致的设备信息格式
        return {
            'mainEquipment': main_equipment,
            'weighingEquipment': weighing_equipment_display_name,
            'weighingEquipmentBaseName': weighing_equipment_base_name,
            'weighingEquipmentId': weighing_equipment_id,
            'weighingEquipmentRaw': weighing_equipments[0]['equipment'] if weighing_equipments else None,
            'titrationEquipment': '',
            'cultivationEquipment': '',
            'mainEquipmentNames': main_equipment_names,
            'mainEquipmentIds': main_equipment_ids,
            'raw_data': equipment_list
        }

    def get_empty_equipment(self):
        """返回空的设备信息"""
        return {
            'mainEquipment': '',
            'weighingEquipment': '',
            'weighingEquipmentId': '',
            'titrationEquipment': '',
            'cultivationEquipment': '',
            'mainEquipmentNames': '',
            'mainEquipmentIds': '',
            'raw_data': []
        }

    def get_main_equipment_choices(self, sample_project_id, log_func=None):
        """查询主检设备可选列表 - 对齐网页端 equipmentBill/ocMultipleChoicePage
        usedCategory=检测设备,主检设备，按 SampleProjectId 查询"""
        return self._query_equipment_choice_page(
            "ocMultipleChoicePage", "检测设备,主检设备",
            sample_project_id, {"Id": "23"}, log_func,  # ponytail: 机构ID=23，与样品列表 decideProjectOrgId 一致
        )

    def get_weighing_equipment_choices(self, sample_project_id, log_func=None):
        """查询称样设备可选列表 - 对齐网页端 equipmentBill/ocChoicePage
        usedCategory=称样设备，按 SampleProjectId 查询"""
        return self._query_equipment_choice_page(
            "ocChoicePage", "称样设备", sample_project_id, {}, log_func,
        )

    def get_project_alias(self, detection_project_id, log_func=None):
        """取项目别名（谱图数据采集的解析规则串）。

        对应网页 检测标准管理→项目→属性 的 detectionProjectProperty/detailByProject，
        返回 (alias, detail)。alias=otherName（如 "Chrysene;<0.005/[0.005-2.00]"）；
        detail 为短状态串（'ok'/'空'/'success=false:...'/'HTTP nnn'/'异常:...'），便于定位。
        id 参数是【检测项目定义ID detectionProjectId】，非样品项目ID。
        """
        try:
            params = {
                'id': detection_project_id,
                'pid': self.get_user_pid(),
                'pname': self.get_user_pname(),
                'loginId': self.get_user_login_id(),
            }
            response = self.login_system.session.get(
                f'{self.login_system.base_url}/detectionManager/manager/detectionProjectProperty/detailByProject',
                params=params,
                headers={
                    'Accept': 'application/json, text/javascript, */*; q=0.01',
                    'Referer': f'{self.login_system.base_url}/web/testStandardMgt.html?menuId=294',
                    'X-Requested-With': 'XMLHttpRequest',
                },
                verify=False, timeout=15,
            )
            if response.status_code == 200:
                result = response.json()
                if result.get('success'):
                    alias = (result.get('resultData') or {}).get('otherName') or ''
                    return alias, ('空' if not alias else 'ok')
                detail = f"success=false: {str(result)[:150]}"
                if log_func:
                    log_func(detail)
                return '', detail
            detail = f"HTTP {response.status_code}"
            if log_func:
                log_func(f"detailByProject {detail}")
            return '', detail
        except Exception as e:
            detail = f"异常: {e}"
            if log_func:
                log_func(f"取项目别名 {detail}")
            return '', detail

    def save_main_equipment(self, experiment_id, items, log_func=None):
        """提交主检设备到 ocExperiment/saveMainEqubment（测试）。
        items: [{label, id, raw}, ...] —— label 形如 "CK-SB107-CG,电热恒温振荡水浴锅,2026-10-10"。
        返回 (success, result_json)。"""
        try:
            entries, ids, name_parts = [], [], []
            for it in items:
                raw = it.get('raw') or {}
                label = it.get('label', '')
                eq_id = it.get('id', '')
                # 编号/名称：优先 raw 字段，回退 label("编号 名称")
                code = (raw.get('no') or raw.get('code') or raw.get('equipmentCode')
                        or raw.get('equipmentNo') or raw.get('number') or raw.get('billCode') or '').strip()
                name = (raw.get('name') or raw.get('equipmentName') or '').strip()
                if not (code and name):
                    sp = label.split(' ', 1)
                    code = code or (sp[0] if sp else '')
                    name = name or (sp[1] if len(sp) > 1 else (sp[0] if sp else ''))
                # 有效日期：取 raw.checkOutDate（前端字段），兼容 "YYYY-MM-DD HH:MM:SS"
                cod = raw.get('checkOutDate')
                check_out = cod[:10] if isinstance(cod, str) else (str(cod)[:10] if cod else '')
                # 仪器条件：selectByDetectionMethodId 的 instrumentConditionArr(JSON 串)→数组，对齐前端 condition
                _icarr = raw.get('instrumentConditionArr')
                try:
                    _cond = json.loads(_icarr) if _icarr else None
                except (ValueError, TypeError):
                    _cond = None
                entries.append({
                    'id': int(eq_id) if str(eq_id).isdigit() else eq_id,
                    'name': name,
                    'checkOutDate': check_out,
                    'condition': _cond,
                })
                _np = ','.join((code, name))
                if check_out:
                    _np += ',' + check_out
                name_parts.append(_np)
                ids.append(str(eq_id))
            request_data = {
                'id': int(experiment_id) if str(experiment_id).isdigit() else 0,
                'mainEquipmentNames': ";".join(name_parts),
                'mainEquipmentIds': ",".join(ids),
                'mainEquipment': json.dumps(entries, ensure_ascii=False),
                'equipmentUseIds': ",".join(['0'] * len(entries)),
                'pid': self.get_user_pid(),
                'pname': self.get_user_pname(),
                'loginId': self.get_user_login_id(),
            }
            if log_func:
                log_func(f"[测试] mainEquipmentNames={request_data['mainEquipmentNames']}")
                log_func(f"[测试] mainEquipmentIds={request_data['mainEquipmentIds']}")
            response = self.login_system.session.post(
                f"{self.login_system.base_url}/detectionManager/manager/ocExperiment/saveMainEqubment",
                json=request_data,
                headers={
                    'Content-Type': 'application/json;charset=UTF-8',
                    'Accept': 'application/json, text/javascript, */*; q=0.01',
                    'Referer': f'{self.login_system.base_url}/web/detectionResultCheckInCalc.html',
                    'X-Requested-With': 'XMLHttpRequest'
                },
                verify=False, timeout=30
            )
            if response.status_code == 200:
                result = response.json()
                return result.get('success', False), result
            return False, {'error': f'HTTP {response.status_code}'}
        except Exception as e:
            if log_func:
                log_func(f"提交主检设备异常: {str(e)}")
            return False, {'error': str(e)}

    def _query_equipment_choice_page(self, action, used_category, sample_project_id, extra_params, log_func=None):
        """通用设备选择页查询，返回 [{label, id}, ...]"""
        try:
            params = {
                "_search": "false",
                "nd": int(time.time() * 1000),
                "pageSize": "30",
                "pageNo": "1",
                "sidx": "",
                "sord": "asc",
                "usedCategory": used_category,
                "SampleProjectId": str(sample_project_id),
                "pid": self.get_user_pid(),
                "pname": self.get_user_pname(),
                "loginId": self.get_user_login_id(),
            }
            params.update(extra_params)
            response = self.login_system.session.get(
                f"{self.login_system.base_url}/detectionManager/manager/equipmentBill/{action}",
                params=params,
                headers={
                    'Accept': 'application/json, text/javascript, */*; q=0.01',
                    'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; WOW64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/101.0.4951.64 Safari/537.36',
                    'Referer': f"{self.login_system.base_url}/web/detectionResultCheckInCalc.html",
                    'X-Requested-With': 'XMLHttpRequest'
                },
                verify=False,
                timeout=30
            )
            if response.status_code == 200:
                result = response.json()
                if result.get('success'):
                    return self._extract_equipment_list(result)
            return []
        except Exception as e:
            if log_func:
                log_func(f"查询设备列表异常({action}): {str(e)}")
            return []

    def _extract_equipment_list(self, result):
        """从设备选择页响应提取设备列表 - 兼容 resultData.voList / rows / list / 平铺列表。
        返回 [{label, id}, ...]：label 含设备编号（mainEquipmentNames 本身为"编号,名称"格式；
        缺失时用 code 字段拼前缀），与 process_equipment_list 的默认设备串可匹配"""
        rd = result.get('resultData')
        if isinstance(rd, dict):
            rows = rd.get('voList') or rd.get('rows') or rd.get('list') or []
        elif isinstance(rd, list):
            rows = rd
        else:
            rows = []
        items, seen = [], set()
        for eq in rows:
            if not isinstance(eq, dict):
                continue
            men = (eq.get('mainEquipmentNames') or '').strip()
            name = (eq.get('name') or eq.get('equipmentName') or '').strip()
            code = (eq.get('no') or eq.get('code') or eq.get('equipmentCode')
                    or eq.get('equipmentNo') or eq.get('number') or eq.get('billCode') or '').strip()
            if men:
                label = men
            elif code and name and code not in name:
                label = f"{code} {name}"
            else:
                label = name
            if not label or label in seen:
                continue
            seen.add(label)
            eq_id = eq.get('equipmentBillId') or eq.get('id') or ''
            items.append({'label': label, 'id': str(eq_id) if eq_id else '', 'raw': eq})
        return items

    def get_current_sample_id(self):
        """获取当前样品ID - 用于构建Referer"""
        return ""

    def get_method_name_by_id(self, method_id, log_func=None):
        """根据方法ID获取方法名称 - 修复API调用"""
        try:
            # 尝试使用不同的API端点
            params = {
                "id": method_id,
                "pid": self.get_user_pid(),
                "pname": self.get_user_pname(),
                "loginId": self.get_user_login_id(),
                "_": str(int(time.time() * 1000))
            }

            # 尝试第一个API端点
            response = self.login_system.session.get(
                f"{self.login_system.base_url}/detectionManager/manager/detectionMethod/getObj",
                params=params,
                headers={
                    'Accept': 'application/json, text/javascript, */*; q=0.01',
                    'Referer': f'{self.login_system.base_url}/web/detectionResultCheckInCalc.html',
                    'X-Requested-With': 'XMLHttpRequest'
                },
                verify=False,
                timeout=30
            )

            if response.status_code == 200:
                result = response.json()
                if result.get('success'):
                    method_data = result.get('resultData', {})
                    method_name = method_data.get('name', '')

                    if method_name:
                        return method_name

            # 如果第一个API失败，尝试第二个API端点
            # 使用项目名称和方法ID查询
            response2 = self.login_system.session.get(
                f"{self.login_system.base_url}/detectionManager/manager/detectionMethod/getObjByIdAndDetProjectName",
                params=params,
                headers={
                    'Accept': 'application/json, text/javascript, */*; q=0.01',
                    'Referer': f'{self.login_system.base_url}/web/detectionResultCheckInCalc.html',
                    'X-Requested-With': 'XMLHttpRequest'
                },
                verify=False,
                timeout=30
            )

            if response2.status_code == 200:
                result = response2.json()
                if result.get('success'):
                    method_list = result.get('resultData', [])
                    for method_data in method_list:
                        if str(method_data.get('decideProjectMethodId')) == str(method_id):
                            method_name = method_data.get('decideProjectMethodName', '')
                            if method_name:
                                return method_name

            return None

        except Exception as e:
            return None

    def _std_no_name_cache_path(self):
        import os
        return os.path.join(os.path.dirname(os.path.abspath(__file__)), ".method_id_cache.json")

    def _load_std_no_name_cache(self):
        import os, json
        try:
            p = self._std_no_name_cache_path()
            if os.path.isfile(p):
                with open(p, encoding="utf-8") as f:
                    return json.load(f)
        except Exception:
            pass
        return {}

    def _save_std_no_name_cache(self):
        import json
        try:
            with open(self._std_no_name_cache_path(), "w", encoding="utf-8") as f:
                json.dump(self._std_no_name_to_id, f, ensure_ascii=False)
        except Exception:
            pass

    def get_method_id_by_standard_no_name(self, standard_no_name, log_func=None):
        """按 standardNoName（如 'AfPS GS 2019:01 PAK 单组份'）解析子方法 methodId。
        一个标准号下常有多个子方法（单组份/N项之和），各为独立 methodId；
        走 selectCheckInDecideMethodName 按 standardNoName 精确匹配取 id。命中返回 id，否则 None。"""
        key = (standard_no_name or "").strip()
        if not key:
            return None
        if key in self._std_no_name_to_id:
            return self._std_no_name_to_id[key]
        try:
            today = datetime.now()
            params = {
                "_search": "false",
                "nd": str(int(time.time() * 1000)),
                "pageSize": "9999",
                "pageNo": "1",
                "sidx": "",
                "sord": "asc",
                "decideProjectMethodName": key,
                "acceptStartDate": (today - timedelta(days=30)).strftime("%Y-%m-%d"),
                "acceptEndDate": today.strftime("%Y-%m-%d"),
                "checkInStatus": "CHECK_IN_STATUS_NO",
                "decideProjectOrgName": "23",
                "subpackage": "NO",
                "pid": self.get_user_pid(),
                "pname": self.get_user_pname(),
                "loginId": self.get_user_login_id(),
            }
            response = self.login_system.session.get(
                f"{self.login_system.base_url}/detectionManager/manager/resultCheckIn/selectCheckInDecideMethodName",
                params=params,
                headers={
                    'Accept': 'application/json, text/javascript, */*; q=0.01',
                    'Referer': f'{self.login_system.base_url}/web/detectionResultCheckInListMgt.html?state=state',
                    'X-Requested-With': 'XMLHttpRequest'
                },
                verify=False,
                timeout=30
            )
            if response.status_code == 200:
                result = response.json()
                if result.get('success'):
                    for m in result.get('resultData', []) or []:
                        if (m.get('standardNoName') or "").strip() == key:
                            mid = m.get('id')
                            self._std_no_name_to_id[key] = mid
                            self._save_std_no_name_cache()
                            return mid
            return None
        except Exception as e:
            if log_func:
                log_func(f"解析子方法ID异常({key}): {str(e)}")
            return None

    def _prefetch_std_no_name_ids(self, names, log_func=None):
        """批量预取子方法ID：一次查30天内全部 decideMethod 记录，按 standardNoName 填缓存。
        替代逐方法名串行查(N方法名=N×RTT)；未命中的名字仍由 get_method_id_by_standard_no_name 单查兜底。"""
        names = {n for n in names if n and n not in self._std_no_name_to_id}
        if not names:
            return
        try:
            today = datetime.now()
            params = {
                "_search": "false", "nd": str(int(time.time() * 1000)),
                "pageSize": "9999", "pageNo": "1", "sidx": "", "sord": "asc",
                "acceptStartDate": (today - timedelta(days=30)).strftime("%Y-%m-%d"),
                "acceptEndDate": today.strftime("%Y-%m-%d"),
                "checkInStatus": "CHECK_IN_STATUS_NO", "decideProjectOrgName": "23",
                "subpackage": "NO", "pid": self.get_user_pid(), "pname": self.get_user_pname(),
                "loginId": self.get_user_login_id(),
            }
            response = self.login_system.session.get(
                f"{self.login_system.base_url}/detectionManager/manager/resultCheckIn/selectCheckInDecideMethodName",
                params=params, headers={
                    'Accept': 'application/json, text/javascript, */*; q=0.01',
                    'Referer': f'{self.login_system.base_url}/web/detectionResultCheckInListMgt.html?state=state',
                    'X-Requested-With': 'XMLHttpRequest'
                }, verify=False, timeout=30)
            if response.status_code == 200:
                result = response.json()
                if result.get('success'):
                    for m in result.get('resultData', []) or []:
                        sn = (m.get('standardNoName') or "").strip()
                        if sn in names:
                            self._std_no_name_to_id[sn] = m.get('id')
        except Exception as e:
            if log_func:
                log_func(f"预取子方法ID异常: {e}")

    def check_and_switch_sub_method(self, method_id, log_func=None):
        """检查并切换子方法 - 修复返回格式问题"""
        if not method_id:
            return method_id

        method_id_str = str(method_id)
        if method_id_str in self.sub_method_map:
            sub_method_info = self.sub_method_map[method_id_str]
            new_method_id = sub_method_info['sub_method_id']
            return new_method_id

        return method_id

    def get_switchable_methods(self, decide_project_name, method_id, log_func=None):
        """查询当前检测项目+方法下可切换的方法ID列表（getObjByIdAndDetProjectName）
        返回 [{decideProjectMethodId, decideProjectMethodName, standardNo, parentId, ...}]"""
        try:
            params = {
                "decideProjectName": decide_project_name,
                "decideProjectMethodId": method_id,
                "pid": self.get_user_pid(),
                "pname": self.get_user_pname(),
                "loginId": self.get_user_login_id(),
                "_": str(int(time.time() * 1000))
            }
            response = self.login_system.session.get(
                f"{self.login_system.base_url}/detectionManager/manager/detectionMethod/getObjByIdAndDetProjectName",
                params=params,
                headers={
                    'Accept': 'application/json, text/javascript, */*; q=0.01',
                    'Referer': f'{self.login_system.base_url}/web/detectionResultCheckInCalc.html',
                    'X-Requested-With': 'XMLHttpRequest'
                },
                verify=False,
                timeout=30
            )
            if response.status_code == 200:
                result = response.json()
                if result.get('success'):
                    return result.get('resultData', [])
            return []
        except Exception as e:
            if log_func:
                log_func(f"查询可切换方法ID异常: {str(e)}")
            return []

    def update_method(self, sample_project_ids, method_id, project_names=None, log_func=None):
        """更新检测方法 - 修复参数格式"""
        try:
            # 构建项目名称字符串 - 用反引号分隔重复的项目名称
            if project_names and len(project_names) > 0:
                # 如果只有一个项目，直接使用项目名称
                if len(project_names) == 1:
                    decide_project_names = project_names[0]
                else:
                    # 多个项目时，用反引号连接相同的项目名称
                    decide_project_names = "`".join([project_names[0]] * len(project_names))
            else:
                decide_project_names = ""

            # 根据抓包数据修正请求参数格式
            request_data = {
                "ids": sample_project_ids,
                "decideProjectMethodId": method_id,
                "decideProjectMethodName": "GCMS法",
                "decideProjectNames": decide_project_names,
                "pid": self.get_user_pid(),
                "pname": self.get_user_pname(),
                "loginId": self.get_user_login_id()
            }

            headers = {
                'Content-Type': 'application/json;charset=UTF-8',
                'Accept': 'application/json, text/javascript, */*; q=0.01',
                'Referer': f'{self.login_system.base_url}/web/detectionResultCheckInCalc.html',
                'X-Requested-With': 'XMLHttpRequest'
            }

            response = self.login_system.session.post(
                f"{self.login_system.base_url}/detectionManager/manager/resultCheckIn/updateMethod",
                json=request_data,
                headers=headers,
                verify=False,
                timeout=30
            )

            if response.status_code == 200:
                result = response.json()
                if result.get('success'):
                    return True
                else:
                    return False
            else:
                return False

        except Exception as e:
            return False

    def get_solution_full_info(self, configure_order, log_func=None):
        """获取标准物质的完整信息 - 简化版本"""
        try:
            # 使用动态查询API获取标准物质信息
            dynamic_params = {
                "_search": "false",
                "nd": int(time.time() * 1000),
                "pageSize": "100",
                "pageNo": "1",
                "sidx": "",
                "sord": "asc",
                "state": "0",
                "keyword": configure_order,
                "type": "SOLUTION_TYPE_D",
                "pid": self.get_user_pid(),
                "pname": self.get_user_pname(),
                "loginId": self.get_user_login_id()
            }

            headers = {
                'Accept': 'application/json, text/javascript, */*; q=0.01',
                'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; WOW64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/101.0.4951.64 Safari/537.36',
                'Referer': f"{self.login_system.base_url}/web/detectionResultCheckInCalc.html",
                'X-Requested-With': 'XMLHttpRequest'
            }

            response = self.login_system.session.get(
                f"{self.login_system.base_url}/detectionManager/manager/ocExperiment/getSolutionAdata",
                params=dynamic_params,
                headers=headers,
                verify=False,
                timeout=30
            )

            if response.status_code == 200:
                result = response.json()
                if result.get('success'):
                    vo_list = result.get('resultData', {}).get('voList', [])
                    # 查找匹配的配置序号
                    for solution in vo_list:
                        if solution.get('configureOrder') == configure_order:
                            # 提取详细信息并构建完整信息字符串
                            solution_name = solution.get('solutionName', '')
                            solution_code = solution.get('solutionCode', '')
                            concentration = solution.get('concentration', '')
                            configure_date = solution.get('configureDate', '')
                            validity_date = solution.get('validityDate', '')

                            full_info = f"{configure_order},{solution_name},{solution_code},{concentration},{configure_date},{validity_date}"
                            return full_info

            # 如果查询失败，返回配置序号
            return configure_order

        except Exception as e:
            if log_func:
                log_func(f"获取标准物质完整信息异常: {str(e)}")
            return configure_order

    def extract_dynamic_columns_from_config(self, experiment_config, log_func=None):
        """从实验配置中提取动态列 - 替代硬编码版本"""
        try:
            dynamic_columns = []

            # 从分析记录中提取动态字段
            analysis_records = experiment_config.get('ocAnalysisRecordList', [])
            if analysis_records:
                # 取第一个记录作为参考
                first_record = analysis_records[0]

                # 查找所有动态字段
                dynamic_fields = {k: v for k, v in first_record.items() if k.startswith('dynamic')}

                for field_code, field_value in dynamic_fields.items():
                    # 从字段代码中提取ID
                    field_id = field_code.replace('dynamic', '')

                    # 尝试从配置中获取字段显示名称
                    display_name = self.get_dynamic_field_display_name_from_system(field_code, field_id, log_func)

                    column = {
                        'id': field_id,
                        'columeCode': field_code,
                        'columeName': display_name,
                        'editType': 'EDIT_TYPE_TEXT',
                        'defaultVal': str(field_value) if field_value else '',
                        'columeOrder': len(dynamic_columns) + 1,
                        'isColumnMerge': 0,  # 兜底列默认不合并（主路径 getDynamicColumns 自带该字段）
                    }
                    dynamic_columns.append(column)

            return dynamic_columns

        except Exception as e:
            if log_func:
                log_func(f"从配置提取动态列异常: {str(e)}")
            return []

    def get_dynamic_field_display_name_from_system(self, field_code, field_id, log_func=None):
        """从系统获取动态字段的显示名称 - 替代硬编码版本"""
        try:
            # 方法1: 尝试通过API获取字段配置
            display_name = self.get_dynamic_field_config(field_id, log_func)
            if display_name:
                return display_name

            # 方法2: 从本地缓存或配置文件中获取
            display_name = self.get_dynamic_field_from_local_config(field_code, log_func)
            if display_name:
                return display_name

            # 方法3: 使用智能推断（基于字段代码模式）
            display_name = self.infer_display_name_from_pattern(field_code, log_func)
            if display_name:
                return display_name

            # 最终回退：返回字段代码
            return field_code

        except Exception as e:
            if log_func:
                log_func(f"获取动态字段显示名称异常: {str(e)}")
            return field_code

    def get_dynamic_field_config(self, field_id, log_func=None):
        """通过API获取动态字段配置"""
        try:
            # 这里可以调用系统的字段配置API
            # 由于系统API可能没有直接提供这个接口，暂时返回None
            return None

        except Exception as e:
            return None

    def get_dynamic_field_from_local_config(self, field_code, log_func=None):
        """从本地配置文件获取动态字段显示名称"""
        try:
            # 可以在这里读取外部配置文件（JSON、YAML等）
            # 示例配置字典 - 可以移到外部文件
            field_mapping = {
                # 这里可以维护一个字段映射表，但不再是硬编码在业务逻辑中
                # 可以定期更新这个映射表
            }

            # 如果配置了映射，返回映射值
            if field_code in field_mapping:
                return field_mapping[field_code]

            return None

        except Exception as e:
            return None

    def infer_display_name_from_pattern(self, field_code, log_func=None):
        """基于字段代码模式智能推断显示名称"""
        try:
            # 常见的字段代码模式推断
            patterns = {
                'dynamic11224': '体积V（ml）',
                'dynamic11225': '样品空白C0（μg/mL）',
                'dynamic11226': '样品浓度C1（μg/mL）',
                'dynamic11227': '稀释因子F',
                'dynamic1692': '体积V（ml）',
                'dynamic1693': '样品空白C0（μg/mL）',
                'dynamic1694': '样品浓度C1（μg/mL）',
                'dynamic1695': '稀释因子F'
            }

            # 精确匹配
            if field_code in patterns:
                return patterns[field_code]

            # 模式匹配（例如：dynamic后跟数字）
            if field_code.startswith('dynamic'):
                # 可以根据字段ID的范围或模式进行推断
                field_id = field_code.replace('dynamic', '')
                if field_id.isdigit():
                    # 这里可以添加更复杂的推断逻辑
                    # 例如：特定范围的ID对应特定类型的字段
                    pass

            return None

        except Exception as e:
            return None


class _Box:
    # ponytail: tk-StringVar/Text 替身，build_grouped_experiment_data 只用 .get()/.set()。
    # remark_text.get('1.0','end-1c') 需接受两参数，故 get 忽略位置参数。仅供无头 host 使用。
    def __init__(self, v=""):
        self._v = v

    def get(self, a=None, b=None):
        return self._v

    def set(self, v):
        self._v = v


def _norm_cn(s):
    """列名/参数名容错归一：NFKC(全角→半角、下标₀→0) + 去全部空白 + 小写。
    覆盖 「分析校正系数e(%)」vs「分析校正系数E（%）」/ C₀ vs C0 等差异，避免固定参数填不上。"""
    return "".join(unicodedata.normalize('NFKC', str(s or '')).split()).lower()


def _match_fixed_params(fixed_params, project):
    """按触发条件筛选适用于该项目的固定参数，返回 {规范化参数名(去空白): 值}。
    触发格式：
      「默认」(无条件，优先级最低，可被具体条件覆盖)；
      「字段=值」单条件；多条件用 ';' 分隔(AND，均需命中)，如「检测项目=可溶性六价铬（CrVI）;标准值=≤0.005」。
    字段：检测项目(projectName)/检测方法(standardNo)/标准值(standardValue)。
    值支持多值(中/英文逗号分隔，任一子串命中即该条件成立)。
    参数名经 _norm_cn 归一(全角/大小写/下标容错)后作为键。"""
    overrides = {}

    def _apply(rule):
        for p in rule.get("params") or []:
            name = _norm_cn(p.get("name"))
            if name:
                overrides[name] = p.get("value", "")

    def _cond_match(cond):
        """单条件 字段=值。值多值逗号 OR；空值=不限(命中)。"""
        field, _, value = cond.partition("=")
        field, value = field.strip(), value.strip()
        if not value:
            return True
        if field == "检测方法":
            target = (project.get("standardNo") or "").strip()
        elif field == "标准值":
            target = (project.get("standardValue") or "").strip()
        else:  # 检测项目(默认)
            target = (project.get("projectName") or "").strip()
        _vals = [v.strip() for v in value.replace("，", ",").split(",") if v.strip()]
        return any(v in target for v in _vals)

    # 先应用默认规则（优先级最低）
    for rule in fixed_params or []:
        if (rule.get("trigger") or "").strip() in ("默认", "默认触发"):
            _apply(rule)
    # 再按条件匹配（覆盖默认）；多条件 ';' AND
    for rule in fixed_params or []:
        trig = (rule.get("trigger") or "").strip()
        if not trig or trig in ("默认", "默认触发"):
            continue
        conditions = [c.strip() for c in trig.replace("；", ";").split(";") if c.strip()]
        if conditions and all(_cond_match(c) for c in conditions):
            _apply(rule)
    return overrides


# ---- 稀释备注 ----
# 稀释.doc：总倍数 F = Π(C/移取量)，C=定容体积；每步因子∈{2,5,10,20}(C=10 时移取 5/2/1/0.5 ml)。
# 13 个标准倍数逐字固化(含 400=[20,20]、2000=[10,10,20] 等非贪婪特例)；表外走贪心分解。
_DILUTION_STEPS = {
    2: [2], 5: [5], 10: [10], 20: [10, 2], 50: [10, 5], 100: [10, 10],
    200: [10, 10, 2], 250: [10, 5, 5], 400: [20, 20], 500: [10, 10, 5],
    1000: [10, 10, 10], 2000: [10, 10, 20], 2500: [10, 10, 5, 5],
}


def _decompose_dilution(f):
    """贪心分解倍数 f 为 [10,5,2] 因子列表(乘积=f)；非 2^a·5^b 返回 None。"""
    if f < 2 or f != int(f):
        return None
    f = int(f)
    steps = []
    for factor in (10, 5, 2):
        while f % factor == 0:
            steps.append(factor)
            f //= factor
    return steps if f == 1 else None


def build_dilution_remark(f, c):
    """按稀释.doc 生成稀释操作备注：移取…样液定容至…,再/最后移取…稀释液定容至….
    f=额外稀释倍数(>1)，c=定容体积。查表优先→贪心分解→单步兜底；无法生成返回 None。
    措辞：步0=样液；末步(n≥3)用「最后移取」、(n==2)用「再移取」；中间步「再移取」，均为稀释液。"""
    try:
        f = float(f); c = float(c)
    except (TypeError, ValueError):
        return None
    if f <= 1 or c <= 0:
        return None
    fi = int(round(f))
    steps = _DILUTION_STEPS.get(fi) or _decompose_dilution(fi) or [fi]  # [fi]=非标准倍数单步兜底
    parts, n = [], len(steps)
    for i, factor in enumerate(steps):
        take = c / factor
        if i == 0:
            parts.append(f"移取{take:.2f}ml样液定容至{c:.2f}ml")
        elif i == n - 1:
            verb = "再移取" if n == 2 else "最后移取"
            parts.append(f"{verb}{take:.2f}ml稀释液定容至{c:.2f}ml")
        else:
            parts.append(f"再移取{take:.2f}ml稀释液定容至{c:.2f}ml")
    return ",".join(parts) + "."


def _we_datetime_str(v):
    """称样设备台账日期 → LIMS 字符串 "YYYY-MM-DD HH:MM:SS"。
    raw_data(selectByDetectionMethodId) 里 createDatetime/modifyDatetime 是 epoch-ms 整型；
    总和无称样列时 LIMS 仅在 weighingEquipment 对象日期可解析时绑定称样设备，整型会致绑定失败。"""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        try:
            return datetime.fromtimestamp(v / 1000.0).strftime("%Y-%m-%d %H:%M:%S")
        except (OSError, ValueError, OverflowError):
            return str(v)
    s = str(v).strip()
    if s.isdigit():
        try:
            return datetime.fromtimestamp(int(s) / 1000.0).strftime("%Y-%m-%d %H:%M:%S")
        except (OSError, ValueError, OverflowError):
            return s
    return s  # 已是字符串日期，原样


def build_grouped_experiment_data(host, projects, experiment_code, method_name, experiment_process_override=None):
    """构建分组实验数据 - 包含完整标准物质信息（共享实现，阶段0从 detection_entry_main 抽出）。

    host 需提供: login_system(.current_user/.users)、experiment_config、equipment_config、
    dynamic_columns、data_fields[columeCode]、remark_text、solution_type_var、
    actual_method_name/actual_method_id(getattr)、api.get_solution_full_info、log、
    temperature_var/humidity_var/start_date_var/end_date_var。
    手动录入传 DetectionEntrySystem 实例；序列无头传 HeadlessHost。
    """
    self = host  # ponytail: 复用原方法体，self 即 host，函数体零改动
    # 获取第一个项目的信息作为基础
    first_project = projects[0]

    # 获取当前用户信息
    current_user = self.login_system.current_user
    user_info = self.login_system.users.get(current_user, {})
    pid = user_info.get('pid', '377')

    try:
        pid_int = int(pid)
    except (ValueError, TypeError):
        pid_int = 377

    # 关键修复：正确构建样品项目ID字符串
    sample_project_ids_list = []
    for project in projects:
        project_id = project.get('projectId')
        if project_id:
            sample_project_ids_list.append(str(project_id))

    sample_project_ids_str = ",".join(sample_project_ids_list)

    # 从动态获取的实验配置中获取关键信息
    oc_method_settings = self.experiment_config.get('ocMethodSettings', {})
    computing_formula = self.experiment_config.get('computingFormula')
    round_method = self.experiment_config.get('roundMethod')
    round_method_level_json = self.experiment_config.get('roundMethodLevelJson')
    calc_method = self.experiment_config.get('calcMethod')
    experiment_process = experiment_process_override or self.experiment_config.get('experimentProcess')

    # 使用保存的实际方法名称和方法ID（可能经过子方法切换）
    actual_method_name = getattr(self, 'actual_method_name', method_name)
    actual_method_id = getattr(self, 'actual_method_id', oc_method_settings.get('methodId'))

    # 获取分析记录配置
    analysis_records_config = self.experiment_config.get('ocAnalysisRecordList', [])

    # 按项目ID分组分析记录
    project_analysis_records = {}
    for record in analysis_records_config:
        project_id = record.get('projectId')
        if project_id not in project_analysis_records:
            project_analysis_records[project_id] = []
        project_analysis_records[project_id].append(record)

    # data_fields 按 ocAnalysisRecordList 位置索引；多组分时第二个组分的记录位置不是 0，
    # 故建 记录对象→全局行号 映射，避免读到第一个组分的输入（串行）
    _record_pos = {id(r): i for i, r in enumerate(analysis_records_config)}

    # 从设备配置中获取设备信息
    main_equipment = self.equipment_config.get('mainEquipment', '')
    main_equipment_names = self.equipment_config.get('mainEquipmentNames', '')
    main_equipment_ids = self.equipment_config.get('mainEquipmentIds', '')
    weighing_equipment = self.equipment_config.get('weighingEquipment', '')
    weighing_equipment_base_name = self.equipment_config.get('weighingEquipmentBaseName', '')
    weighing_equipment_id = self.equipment_config.get('weighingEquipmentId', '')

    # 修复设备信息格式 - 与前端 saveOcExperiment 一致：发完整台账对象(id/name + 元数据)。
    # 无称样列的实验(如总和) LIMS 仅在收到完整对象时绑定称样设备(单组份有称样列,{id,name} 即可)
    _we_raw = self.equipment_config.get('weighingEquipmentRaw')
    if weighing_equipment_id:
        weighing_equipment_obj = {
            "id": weighing_equipment_id,
            "name": weighing_equipment_base_name or weighing_equipment,
        }
        if isinstance(_we_raw, dict):
            for _k in ("creatorName", "creatorId", "modifierName", "createDatetime", "modifyDatetime"):
                _v = _we_raw.get(_k)
                if _v is None:
                    continue
                if _k in ("createDatetime", "modifyDatetime"):
                    _v = _we_datetime_str(_v)
                weighing_equipment_obj[_k] = _v
    else:
        weighing_equipment_obj = ""

    weighing_equipment_json = weighing_equipment

    # 构建主检设备数组格式 - 与前端保持一致
    # 仪器条件 condition：selectByDetectionMethodId 的 instrumentConditionArr(JSON 串)按 id 回填，对齐前端
    _cond_by_id = {}
    for _eq in (self.equipment_config.get("raw_data") or []):
        _icarr = _eq.get("instrumentConditionArr")
        if not _icarr:
            continue
        for _kid in (_eq.get("equipmentBillId"), _eq.get("id")):
            if _kid is not None:
                try:
                    _cond_by_id[str(_kid)] = json.loads(_icarr)
                except (ValueError, TypeError):
                    pass
    main_equipment_array = []
    if main_equipment_ids:
        equipment_ids = main_equipment_ids.split(',')
        equipment_names = main_equipment.split('; ')

        for i, equipment_id in enumerate(equipment_ids):
            equipment_name = ""
            equipment_checkout_date = ""

            if i < len(equipment_names):
                full_name = equipment_names[i]
                # 解析设备名称和日期信息
                if ',' in full_name:
                    parts = full_name.split(',')
                    if len(parts) >= 2:
                        equipment_name = parts[1].strip()
                    if len(parts) >= 3:
                        equipment_checkout_date = parts[2].strip()
                else:
                    equipment_name = full_name

                # 构建与前端完全一致的设备对象结构
                equipment_info = {
                    "id": equipment_id,
                    "name": equipment_name,
                    "checkOutDate": equipment_checkout_date,
                    "condition": _cond_by_id.get(str(equipment_id)),
                    "instrumentAttachId": None
                }
                main_equipment_array.append(equipment_info)

    # 默认触发的固定参数同时写死「计算值+报告值」时，结果全固定，无需计算公式/修约/计算方法，跳过后端 calcTheValue
    _default_fp = next((r for r in (getattr(self, "fixed_params", None) or [])
                        if (r.get("trigger") or "").strip() in ("默认", "默认触发")), None)
    _default_names = {_norm_cn(p.get("name"))
                      for p in ((_default_fp or {}).get("params") or [])}
    skip_calc = _default_names >= {"计算值", "报告值"}

    # 检查关键配置是否存在（固定结果值的方法可缺计算公式/修约/计算方法）
    if not skip_calc:
        if not computing_formula:
            raise Exception("实验配置中缺少计算公式")
        if not round_method:
            raise Exception("实验配置中缺少修约方法")
        if not calc_method:
            raise Exception("实验配置中缺少计算方法")
    if not experiment_process:
        raise Exception("实验配置中缺少实验过程")

    # 修复：将 calcMethod 从对象转换为字符串
    if isinstance(calc_method, dict):
        calc_method_str = calc_method.get('key', 'CALC_METHOD_ENUM_AVG')
    elif calc_method:
        calc_method_str = str(calc_method)
    else:
        calc_method_str = None  # 固定结果值方法无 calcMethod，提交 null 而非 "None" 字符串

    # 修复：确保 ocMethodSettings 中的枚举字段是字符串而不是对象
    if oc_method_settings and isinstance(oc_method_settings, dict):
        cleaned_oc_method_settings = {}
        for key, value in oc_method_settings.items():
            if key in ['id', 'curveCalcMethod', 'curveXCalc', 'curveYCalc', 'curveCalcXyReverse',
                       'isCalculatedValue', 'sampleEvaluate', 'reportUnit', 'isOptionalTable',
                       'calculatedValueWidth', 'calculatedUnitWidth', 'reportValueWidth',
                       'reportUnitWidth', 'isRepeatabilityLimit', 'associatedRecords']:
                if key in ['curveCalcMethod', 'curveXCalc', 'curveYCalc']:
                    if isinstance(value, dict) and 'key' in value:
                        cleaned_oc_method_settings[key] = value['key']
                    else:
                        cleaned_oc_method_settings[key] = value
                else:
                    cleaned_oc_method_settings[key] = value
        oc_method_settings = cleaned_oc_method_settings

    # 构建分析记录列表
    oc_analysis_record_save_list = []
    calc_rows = []  # calcTheValue 请求行，与分析记录共用 id
    record_id_counter = 1

    # 获取固定备注字段的内容
    remark_content = ""
    if hasattr(self, 'remark_text'):
        remark_content = self.remark_text.get('1.0', 'end-1c').strip()

    # 稀释备注：解析稀释列与定容体积列码（启用稀释备注时；列名无关，按 equipRelativeTitle 找稀释列）
    _dil_remark_on = getattr(self, "dilution_remark_enabled", False)
    _dil_col_code = next((c.get('columeCode') for c in self.dynamic_columns
                          if isinstance(c, dict) and (c.get('equipRelativeTitle') or '').strip() == '稀释'), None)
    _vol_col_code = None
    if _dil_remark_on and _dil_col_code:
        _vol_want = _norm_cn(getattr(self, "dilution_volume_column", "") or "")
        if _vol_want:
            for c in self.dynamic_columns:
                if isinstance(c, dict) and _norm_cn(c.get('columeName', '')) == _vol_want:
                    _vol_col_code = c.get('columeCode'); break
            if not _vol_col_code:  # 兜底双向子串
                for c in self.dynamic_columns:
                    _cn_v = _norm_cn(c.get('columeName', '')) if isinstance(c, dict) else ''
                    if _cn_v and (_vol_want in _cn_v or _cn_v in _vol_want):
                        _vol_col_code = c.get('columeCode'); break
            if not _vol_col_code:
                self.log(f"稀释备注: 找不到定容体积列「{getattr(self, 'dilution_volume_column', '')}」，实际列: "
                         f"{[c.get('columeName') for c in self.dynamic_columns if isinstance(c, dict)]}")

    for i, project in enumerate(projects):
        sample_id = project.get('sampleId')
        project_id = project.get('projectId')
        detection_no = project.get('detectionNo')
        sample_small_no = project.get('sampleSmallNo')
        sample_name = project.get('sampleName', '未知样品')

        # 固定参数覆盖（序列模式：方法 other_params_settings.fixed_params，按触发条件匹配当前项目）
        # 附上 standardValue(从分析记录配置)，供触发条件「标准值=...」匹配
        _rec_sv = next((r.get('standardValue') for r in (self.experiment_config.get('ocAnalysisRecordList') or [])
                        if str(r.get('projectId')) == str(project_id)), None)
        _p_for_match = dict(project)
        _p_for_match['standardValue'] = _rec_sv or ''
        fixed_overrides = _match_fixed_params(getattr(self, "fixed_params", None), _p_for_match)
        if fixed_overrides:
            self.log(f"[固定参数] 项目={project.get('projectName')!r} 命中触发条件，待填: {list(fixed_overrides)}")
            _col_norm = {_norm_cn(c.get('columeName')) for c in self.dynamic_columns if isinstance(c, dict)}
            _miss = [k for k in fixed_overrides if k not in _col_norm]
            if _miss:
                self.log(f"[固定参数] 找不到同名列(参数名↔列名不符)，未填: {_miss}")
                self.log(f"[固定参数] 实际动态列名: {[c.get('columeName') for c in self.dynamic_columns if isinstance(c, dict)]}")

        # 保留名「计算值/报告值」：直接固定结果值，跳过后端 calcTheValue（按方法触发，该项目下所有记录统一）
        fixed_calc_value = fixed_overrides.pop("计算值", None)
        fixed_report_value = fixed_overrides.pop("报告值", None)

        # 检查必要字段
        if not sample_id:
            raise Exception(f"项目 {project.get('projectName')} 缺少样品ID")
        if not project_id:
            raise Exception(f"项目 {project.get('projectName')} 缺少项目ID")
        if not detection_no:
            raise Exception(f"项目 {project.get('projectName')} 缺少检测编号")
        if not sample_small_no:
            raise Exception(f"项目 {project.get('projectName')} 缺少样品小号")

        # 获取该项目对应的分析记录
        project_records = project_analysis_records.get(project_id, [])

        if not project_records:
            project_records = [{
                "serialNumber": 1,
                "reportUnitName": "mg/kg",
                "calculatedUnit": "mg/kg",
                "detectionLimit": "0.1",
                "standardValue": "≤50",
                "detectionLimitType": "检出限"
            }]

        # 修正 detectionNoAndSmallNo 格式 - 与前端保持一致
        detection_no_and_small_no = f"{detection_no}-{sample_small_no.zfill(3)}"

        # 为每个分析记录生成记录
        for record_index, record_config in enumerate(project_records):
            # 动态构建字段数据
            dynamic_fields = {}
            selectmap_data = {}
            dil_extra = 1.0  # 额外稀释倍数(报告解析/文件名扫描)；启用稀释备注时用于生成备注

            # 使用用户输入的值覆盖默认值
            _is_headless = getattr(self, "is_headless", False)
            for col in self.dynamic_columns:
                col_id = col.get('id')
                col_code = col.get('columeCode', f'dynamic{col_id}')
                col_name = col.get('columeName', '')
                edit_type = col.get('editType', 'EDIT_TYPE_TEXT')
                default_val = col.get('defaultVal', '')

                # 从用户输入获取值（合并列共享一个值；非合并列按测试次数取对应行）
                user_value = ""
                if hasattr(self, 'data_fields') and col_code in self.data_fields:
                    field = self.data_fields[col_code]
                    if isinstance(field, list):
                        _gi = _record_pos.get(id(record_config), record_index)
                        if 0 <= _gi < len(field):
                            user_value = field[_gi].get().strip()
                    else:
                        user_value = field.get().strip()

                # 如果用户输入了值，使用用户输入的值；否则使用配置中的默认值
                if user_value:
                    dynamic_fields[col_code] = user_value
                else:
                    # 如果用户没有输入，使用配置记录中的值或默认值
                    config_value = record_config.get(col_code)
                    if config_value is not None:
                        dynamic_fields[col_code] = str(config_value) if config_value is not None else default_val
                    else:
                        dynamic_fields[col_code] = default_val

                # 固定参数：仅当该列无真实值(手填或报告解析填入)时补全；有真实值则尊重不覆盖。
                # 序列(无头)模式 data_fields 预填的是列 defaultVal(占位)，并非真实输入——
                # 占位值(==defaultVal)应被固定参数覆盖；但报告解析填入的实测浓度(如「样品浓度C」，≠defaultVal)
                # 不被覆盖(否则可溶性汞等由 样品浓度C 计算的结果会被钉死在 <0.004)。
                _cn = _norm_cn(col_name)
                _placeholder = _is_headless and user_value == str(default_val or "").strip()
                if _cn and _cn in fixed_overrides and (not user_value or _placeholder):
                    dynamic_fields[col_code] = str(fixed_overrides[_cn])

                # 稀释列(启用稀释备注时)：最终值=基准(固定参数稀释因子F)×额外(报告解析/文件名扫描)，
                # 额外倍数留给备注。基准不在固定参数时 _b=1(列值即额外，与现状一致)。
                if _dil_remark_on and _dil_col_code and col_code == _dil_col_code:
                    try:
                        _e = float(user_value or default_val or 1)
                    except (ValueError, TypeError):
                        _e = 1.0
                    _b = 1.0
                    if _cn in fixed_overrides:
                        try:
                            _b = float(fixed_overrides[_cn] or 1)
                        except (ValueError, TypeError):
                            _b = 1.0
                    dil_extra = _e
                    dynamic_fields[col_code] = f"{_b * _e:g}"

                # 构建selectmap数据 - 与前端保持一致
                if edit_type == 'EDIT_TYPE_SELECT' and dynamic_fields[col_code]:
                    selectmap_data[col_code] = [{
                        "id": col_id,
                        "createDatetime": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                        "methodSettingsId": oc_method_settings.get('id', ''),
                        "methodTitleSettingId": col_id,
                        "sort": 1,
                        "isDefault": 1,
                        "selectValueName": dynamic_fields[col_code],
                        "creatorName": "系统管理员",
                        "creatorId": 1,
                        "modifierName": "系统管理员",
                        "modifyDatetime": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                    }]

            # 从分析记录配置中获取关键参数
            detection_limit = record_config.get('detectionLimit', '0.1')
            standard_value = record_config.get('standardValue', '≤1.0')
            standard_and_warning_val = record_config.get('standardAndWarningVal', '≤1.0')
            detection_limit_type = record_config.get('detectionLimitType', '检出限')
            accuracy = record_config.get('accuracy', 'STANDARD_DEVIATION')
            # 稀释备注(按额外倍数+定容体积)：仅额外>1 的稀释样品生成，否则沿用固定备注
            row_remark = remark_content
            if _dil_remark_on and _vol_col_code and dil_extra > 1:
                try:
                    _c = float(dynamic_fields.get(_vol_col_code) or 0)
                except (ValueError, TypeError):
                    _c = 0.0
                if _c > 0:
                    row_remark = build_dilution_remark(dil_extra, _c) or remark_content
            sample_remark_content = row_remark
            report_unit_name = record_config.get('reportUnitName', 'mg/kg')
            calculated_unit = record_config.get('calculatedUnit', 'mg/kg')

            # 构建分析记录 - 与前端数据结构完全一致
            analysis_record = {
                "id": record_id_counter,
                "createDatetime": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                "ocExperiment": None,
                "serialNumber": record_config.get('serialNumber', record_index + 1),
                "sampleId": sample_id,
                "sampleCode": detection_no + sample_small_no,
                "sampleName": sample_name,
                "projectId": project_id,
                "projectCode": None,
                "projectName": project.get('projectName', ''),
                "originalNumber": None,
                "detectionLimit": detection_limit,
                "standardValue": standard_value,
                "reportValue": None,
                "reportUnit": None,
                "reportUnitName": report_unit_name,
                "other": None,
                "spectrumId": None,
                "description": None,
                "calculatedValue": None,
                "calculatedUnit": calculated_unit,
                "accuracy": accuracy,
                "remark": row_remark,
                "precisionCalcEnumMap": "{\"PROCESS_STATUS_NEEDLESS\":\"不需要\",\"STANDARD_DEVIATION\":\"标准偏差\",\"RELATIVE_LABEL_DEVIATION\":\"相对标准偏差(>2)\",\"RELATIVE_LABEL_DEVIATION_TWO\":\"相对标准偏差(≥2)\",\"RELATIVE_DIFFERENCE\":\"相对相差(误差)\",\"ABSOLUTE_DEVIATION\":\"绝对偏差\",\"RELATIVE_DEVIATION\":\"相对偏差\",\"ABSOLUTE_DIFFERENCE\":\"绝对差值\"}",
                "ocAnalysisExperimentList": None,
                "ocCurve": None,
                "ocCurveSave": None,
                "ocCurveValueSaveList": None,
                "sampleProjectId": project_id,
                "detectionLimitType": detection_limit_type,
                "originalNo": None,
                "standardAndWarningVal": standard_and_warning_val,
                "creatorName": None,
                "creatorId": None,
                "modifierName": None,
                "modifyDatetime": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                # 动态字段
                **dynamic_fields,
                # 添加selectmap字段
                "selectmap": json.dumps(selectmap_data, ensure_ascii=False) if selectmap_data else "{}",
                "detectionNo": detection_no,
                "detectionNoAndSmallNo": detection_no_and_small_no,
                "samplePhotos": False,
                "sampleEvaluate": None,
                "sampleRemarkContent": sample_remark_content,
                "allSelectStatus": False,
                "picColorStatus": False,
                "cid": None,
                "_X_ID": f"row_{8 + record_id_counter}"
            }
            # 固定结果值写回（来自固定参数保留名「计算值/报告值」）
            if fixed_calc_value is not None:
                analysis_record["calculatedValue"] = fixed_calc_value
            if fixed_report_value is not None:
                analysis_record["reportValue"] = fixed_report_value
            oc_analysis_record_save_list.append(analysis_record)
            # 收集 calcTheValue 请求行（columnValues=动态字段+检出限，calculatedValue 留空由服务端算）
            calc_rows.append({
                "columnValues": {
                    "detectionLimit": detection_limit,
                    "detectionLimitType": detection_limit_type,
                    # 空值发 null 对齐网页端；发 "" 服务端按数字解析失败会返回空 calculatedValue
                    **{k: (v if v != "" else None) for k, v in dynamic_fields.items()},
                    "calculatedValue": "",
                },
                "id": record_id_counter,
                "projectId": project_id,
                "sampleCode": detection_no + sample_small_no,
                "calcUnit": calculated_unit,
                "reportUnit": report_unit_name,
                "serialNum": record_config.get('serialNumber', record_index + 1),
            })
            record_id_counter += 1

    # 计算报告值并回填（calcTheValue）；固定结果值时跳过；失败回退原 None 行为，不阻断提交
    calc_result = None
    if not skip_calc:
        calc_meta = {
            'roundMethod': round_method,
            'roundMethodLevelJson': round_method_level_json,
            'resultRoundMethod': self.experiment_config.get('resultRoundMethod'),
            'resultRoundMethodLevelJson': self.experiment_config.get('resultRoundMethodLevelJson'),
            'calcMethod': calc_method_str,
            'accuracy': (analysis_records_config[0].get('accuracy') if analysis_records_config else None) or 'STANDARD_DEVIATION',
            'methodSettingId': (self.experiment_config or {}).get('ocMethodSettings', {}).get('methodId') or oc_method_settings.get('id'),
            'accuracyRadixPoint': oc_method_settings.get('accuracyRadixPoint'),
            'curve': self.experiment_config.get('ocCurve'),
        }
        calc_result = self.api.calc_report_values(calc_rows, calc_meta, self.log)
        if calc_result:
            for rec in oc_analysis_record_save_list:
                cv = calc_result.get(rec.get('id'))
                if cv:
                    # 已被固定参数写死的值不覆盖
                    if rec.get('calculatedValue') is None:
                        rec['calculatedValue'] = cv.get('calculatedValue')
                    if rec.get('reportValue') is None:
                        rec['reportValue'] = cv.get('reportValue')
                    if cv.get('other'):
                        rec['other'] = cv['other']
            if self.log:
                self.log(f"已计算报告值: {len(calc_result)} 条")
    elif self.log:
        self.log("结果值由固定参数写死，跳过 calcTheValue")

    # 构建检测方法对象 - 与前端保持一致
    detection_method = {
        "standardNo": actual_method_name,
        "id": actual_method_id if actual_method_id else "",
        "subMethodName": None
    }

    # 获取标准物质信息（支持逗号分隔多个标液，逐个取完整信息后拼接）
    reference_material = ""
    if hasattr(self, 'solution_type_var') and self.solution_type_var.get().strip():
        _std_orders = [o.strip() for o in self.solution_type_var.get().replace("，", ",").split(",") if o.strip()]
        reference_material = "; ".join(self.api.get_solution_full_info(o, self.log) for o in _std_orders)

    # 构建基础实验数据 - 与前端数据结构完全一致
    experiment_data = {
        "sampleProjectIds2": "",
        "newSampleProjectIds": "",
        "experimentCode": experiment_code,
        "fileCode": first_project.get('detectionNo', 'TR25100009'),
        # 使用完整的标准物质信息
        "referenceMaterial": reference_material,
        "standardStrain": None,
        "reagent": None,
        "cultureMedium": None,
        "temperature": self.temperature_var.get(),
        "humidity": self.humidity_var.get(),
        "detectionMethod": detection_method,
        "computingFormula": computing_formula,
        "experimentProcess": experiment_process,
        "methodDescription": None,
        "weighingEquipment": weighing_equipment_obj,
        "weighingEquipmentJson": weighing_equipment_json,
        "cultivationEquipmentJson": "",
        "titrationEquipmentJson": "",
        "cultivationEquipment": "",
        "titrationEquipment": "",
        "reagentId": "",
        "titrationSolution": None,
        "userId": "",
        "description": None,
        "otherDescription": None,
        "mainEquipmentNames": main_equipment_names,
        "detectionMethodEquipmentBills": [],
        "mainEquipmentIds": main_equipment_ids,
        "ocCurve": {
            "id": "",
            "equation": "",
            "correlationCoefficient": "",
            "slope": "",
            "intercept": ""
        },
        "roundMethod": round_method,
        "roundMethodLevelJson": round_method_level_json,
        "calcMethod": calc_method_str,
        "ocCurveSaveList": "[]",
        "fileIds": "",
        "ocMethodSettings": oc_method_settings,
        "spectrumJsonList": "",
        "recoverySpectrumJsonList": "",
        "consumableReceiveIds": None,
        "titrationConsumableReceiveIds": None,
        "emptyColonyCountList": None,
        "compareColonyCountList": None,
        "selfColonyCountList": None,
        "emptySampleColonyCountList": None,
        "negativeColonyCountList": None,
        "keepSampleColonyCountList": None,
        "curveCheckColonyCountList": None,
        "jcxTestColonyCountList": None,
        "materialColonyCountList": None,
        "methodCheckCode": None,
        "resultRoundMethod": None,
        "resultRoundMethodLevelJson": None,
        "mainEquipment": json.dumps(main_equipment_array, ensure_ascii=False) if main_equipment_array else "[]",
        "dynamicColumns": json.dumps(self.dynamic_columns, ensure_ascii=False) if self.dynamic_columns else "[]",
        "startTime": self.start_date_var.get() + " 00:00:00",
        "endTime": self.end_date_var.get() + " 00:00:00",
        "KeyTime": "",
        "ocAnalysisRecordSaveList": json.dumps(oc_analysis_record_save_list, ensure_ascii=False),
        "configureJsonList": "[]",
        "configureDJsonList": "[]",
        "configureSJsonList": "[]",
        "selfRecoverySaveList": "[]",
        "emptyRecoverySaveList": "[]",
        "negativeRecoverySaveList": "[]",
        "npecompareRecoverySaveList": "[]",
        "materialRecoverySaveList": "[]",
        "curveCheckRecoverySaveList": "[]",
        "emptyTestRecoverySaveList": "[]",
        "keepSampleRecoverySaveList": "[]",
        "jcxTestRecoverySaveList": "[]",
        # 关键修复：使用所有项目的ID
        "sampleProjectIds": sample_project_ids_str,
        "pid": pid_int,
        "pname": current_user,
        "loginId": pid_int
    }

    return experiment_data


if __name__ == "__main__":
    # generate_experiment_code 唯一性自检：同秒连发多批不应撞号(XRF 多方法各一批曾因此合并)
    _api = object.__new__(DetectionAPI)
    _api.experiment_code_cache = {}
    _api.user_prefix_cache = {}
    _api.get_user_pname = lambda: "lqy"
    _codes = [_api.generate_experiment_code(method_name=f"m{i}", force_new=False) for i in range(4)]
    assert len(set(_codes)) == 4, _codes                       # 同秒4批各不相同
    assert all(c.startswith("lqy") and len(c) == 17 for c in _codes), _codes  # 格式不变
    assert _api.generate_experiment_code(method_name="m0") == _codes[0]       # 同方法非force_new 复用缓存
    # 称样设备台账日期归一化自检：整型 epoch-ms → "YYYY-MM-DD HH:MM:SS" 字符串(总和绑定称样设备依赖此)
    assert _we_datetime_str(1772507451000).startswith("20") and isinstance(_we_datetime_str(1772507451000), str)
    assert _we_datetime_str("2026-05-07 09:00:44") == "2026-05-07 09:00:44"  # 已是字符串，原样
    assert _we_datetime_str(None) is None                                     # None 不填该字段
    # _match_fixed_params 多值触发自检：中/英文逗号分隔，任一子串命中即应用
    _fp = [{"trigger": "检测项目=可溶性铅（Pb），可溶性镉（Cd），可溶性汞（Hg）",
            "params": [{"name": "称样量m(g)", "value": "0.5000"}]}]
    assert _match_fixed_params(_fp, {"projectName": "可溶性铅（Pb）"}) == {"称样量m(g)": "0.5000"}
    assert _match_fixed_params(_fp, {"projectName": "可溶性汞（Hg）"}) == {"称样量m(g)": "0.5000"}
    assert _match_fixed_params(_fp, {"projectName": "苯"}) == {}                 # 不在列表→不命中
    # 多条件 ';' AND + 标准值：同名项目按标准值区分
    _fp2 = [{"trigger": "检测项目=可溶性六价铬（CrVI）;标准值=≤0.005",
             "params": [{"name": "稀释因子F", "value": "2.5"}]}]
    _nk = _norm_cn("稀释因子F")
    assert _match_fixed_params(_fp2, {"projectName": "可溶性六价铬（CrVI）", "standardValue": "≤0.005"})[_nk] == "2.5"
    assert _match_fixed_params(_fp2, {"projectName": "可溶性六价铬（CrVI）", "standardValue": "≤0.01"}) == {}  # 标准值不符
    # 全角分号；与半角;等价(用户易输入全角)
    _fp3 = [{"trigger": "检测项目=可溶性六价铬（CrVI）；标准值=≤0.005",
             "params": [{"name": "稀释因子F", "value": "2.5"}]}]
    assert _match_fixed_params(_fp3, {"projectName": "可溶性六价铬（CrVI）", "standardValue": "≤0.005"})[_nk] == "2.5"
    # build_dilution_remark 自检：稀释.doc 写法(总倍数 F=Π(C/移取量))
    assert build_dilution_remark(20, 10) == "移取1.00ml样液定容至10.00ml,再移取5.00ml稀释液定容至10.00ml."
    assert build_dilution_remark(10, 25) == "移取2.50ml样液定容至25.00ml."        # C 缩放
    assert build_dilution_remark(400, 10) == "移取0.50ml样液定容至10.00ml,再移取0.50ml稀释液定容至10.00ml."
    assert build_dilution_remark(1000, 10).count("最后移取") == 1                  # 3 步末步「最后移取」
    assert build_dilution_remark(25, 10) == "移取2.00ml样液定容至10.00ml,再移取2.00ml稀释液定容至10.00ml."  # 非标准→分解[5,5]
    assert build_dilution_remark(1, 10) is None and build_dilution_remark(5, 0) is None  # 不稀释/无体积→None
    print("detection_entry_api selfcheck OK")